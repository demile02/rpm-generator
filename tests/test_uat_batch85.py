#!/usr/bin/env python3
"""Batch 8.5: UAT content & structure fixes — targeted regression tests.

1. readiness dari input guru (tanpa diagnosis AI)
2. alasan relevansi dimensi
3. hapus blok standalone KBC
4. KKTP tanpa metadata internal
5. refinement topik
6. prinsip per pertemuan
7. jumlah soal asesmen
8. LKPD tidak dirender/default
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from module_generation_pipeline import (
    MasterOutlineValidator, ModuleGenerationPipeline,
)

MEET_PARAMS = {
    'education_system': 'KEMENAG',
    'institution_type': 'MTs',
    'grade': 'MTs_7',
    'subject': 'Akidah Akhlak',
    'element': 'Pemahaman Konsep',
    'topic': 'Zakat dan wakaf',
    'requested_tp_count': 3,
    'jumlah_pertemuan': 2,
    'jp_per_pertemuan': 2,
    'penyusun': 'Ibu Guru',
    'student_readiness': 'Sebagian murid sudah memahami istilah Zakat, '
                         'namun sebagian besar belum bisa mengaplikasikan '
                         'perhitungan zakat ke dalam kehidupan sehari-hari.',
}

ROWS_D = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-1', 'mengaplikasi'),
    (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'memahami'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
]


@pytest.fixture
def pipeline(test_db):
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


def _mock_activities():
    from test_phase_c_module_generation import mock_meeting_principles
    return {'activities': [
        {'name': f'A{i + 1}', 'description': 'Deskripsi ' + f'A{i + 1}',
         'duration': dur, 'experience': exp, 'tp_linked': tp,
         'stage': stage, 'meeting': meeting}
        for i, (meeting, stage, dur, tp, exp) in enumerate(ROWS_D)],
        'meeting_principles': mock_meeting_principles(2)}


def _generate(pipeline, params=None, assessments=None):
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            return _mock_activities()
        return fake_router_response(prompt, system_instruction, **kwargs)

    with patch_pipeline_ai(pipeline, side_effect):
        return pipeline.generate(dict(params or MEET_PARAMS))


def _docx_text(module_dict):
    import tempfile
    import os
    from docx import Document
    from docx_renderer import render_final_rpm_docx
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    try:
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_dict,
             'validation': module_dict.get('validation_results', {})},
            path)
        doc = Document(out)
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.extend(c.text for c in row.cells)
        return '\n'.join(texts)
    finally:
        os.unlink(path)


class TestReadinessPreservation:
    """§1: ringkasan = teks guru; tanpa diagnosis AI."""

    def test_summary_is_teacher_text(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        readiness = result['module']['master_outline'][
            'identification']['learner_readiness']
        assert readiness['summary'] == MEET_PARAMS['student_readiness']
        assert readiness['aspects'] == []
        text = _docx_text(result['module'])
        assert MEET_PARAMS['student_readiness'][:40] in text

    def test_no_invented_diagnosis_terms(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        readiness = result['module']['master_outline'][
            'identification']['learner_readiness']
        blob = (readiness['summary'] + ' ' +
                ' '.join(readiness['aspects'])).lower()
        for invented in ('nisab', 'haul', '2,5', 'infak', 'sedekah', 'wakaf'):
            if invented not in MEET_PARAMS['student_readiness'].lower():
                assert invented not in blob, invented

    def test_empty_input_stays_honest(self, pipeline):
        params = dict(MEET_PARAMS)
        params.pop('student_readiness')
        result = _generate(pipeline, params)
        assert result['status'] == 'success', result['errors']
        readiness = result['module']['master_outline'][
            'identification']['learner_readiness']
        assert readiness['summary'] == ''
        assert 'belum tersedia' in readiness['note']


class TestDimensionNotes:
    """§2: setiap dimensi terpilih punya alasan relevansi."""

    def test_notes_present_and_specific(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        notes = module['profile_dimension_notes']
        for dim in module['profile_dimensions']:
            note = notes.get(dim)
            assert isinstance(note, str) and len(note.strip()) >= 20, dim
            assert note.strip().lower() != dim.strip().lower()
        text = _docx_text(module)
        assert 'Dimensi Profil Lulusan' in text
        assert notes[module['profile_dimensions'][0]][:30] in text

    def test_missing_note_fails_closed(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def no_notes(prompt, system_instruction=None, **kwargs):
            resp = fake_router_response(prompt, system_instruction, **kwargs)
            if 'triggering_questions' in prompt:
                resp = dict(resp)
                resp.pop('profile_dimension_notes', None)
            return resp

        with patch_pipeline_ai(pipeline, no_notes):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'failed'
        assert result['module'] is None


class TestKbcStandaloneRemoved:
    """§3: tanpa blok 'Penerapan dalam pembelajaran'; traceability utuh."""

    def test_no_standalone_block_docx(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'Penerapan dalam pembelajaran' not in text
        assert 'Kurikulum Berbasis Cinta' in text

    def test_no_standalone_block_but_data_kept(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        # Data internal tetap ada (traceability), hanya tidak dirender.
        assert kbc.get('integration_note')
        assert kbc['themes']

    def test_traceability_intact(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        assert kbc['enabled'] is True
        assert kbc['themes']
        assert kbc.get('integration_note')


class TestKktpPlain:
    """§4: KKTP-Y saja; linkage tetap di backend."""

    def test_docx_no_internal_metadata(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'KKTP-1' in text
        assert '(untuk TP-1' not in text
        assert 'level C4' not in text

    def test_linkage_still_validated(self, pipeline):
        from test_factuality_gate import module_from_result
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = module_from_result(result)
        ok, errors = MasterOutlineValidator.validate(module)
        assert ok, errors
        tp_ids = {t.id for t in module.learning_objectives}
        assert all(k.tp_id in tp_ids for k in module.success_criteria)


class TestTopicRefinement:
    """§5: seed -> topik formal; substansi terjaga; tak selalu verbatim."""

    def test_refined_differs_but_related(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        assert module['topic_seed'] == 'Zakat dan wakaf'
        assert module['topic'] == 'Zakat dan wakaf dalam Pembelajaran'
        assert module['topic'] != module['topic_seed']
        assert 'Zakat dan wakaf' in module['title']

    def test_drift_rejected(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def drifting(prompt, system_instruction=None, **kwargs):
            if 'topik_rpm' in prompt:
                return {'topik_rpm': 'Sejarah Kerajaan Majapahit'}
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(pipeline, drifting):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_verbatim_echo_rejected(self, pipeline):
        """UAT FIX: echo mentah input guru wajib ditolak (fail-closed)."""
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def echo(prompt, system_instruction=None, **kwargs):
            if 'topik_rpm' in prompt:
                return {'topik_rpm': 'Zakat dan wakaf'}
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(pipeline, echo):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_whitespace_dims_still_paired(self, pipeline):
        """UAT FIX: 'Kolaborasi ' (spasi) tetap terpasang ke catatannya."""
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def spaced(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _mock_activities()
            resp = fake_router_response(prompt, system_instruction, **kwargs)
            if 'triggering_questions' in prompt:
                resp = dict(resp)
                resp['profile_dimensions'] = ['Kolaborasi ', 'Kreativitas']
            return resp

        with patch_pipeline_ai(pipeline, spaced):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'success', result['errors']
        module = result['module']
        assert module['profile_dimensions'] == ['Kolaborasi', 'Kreativitas']
        assert set(module['profile_dimension_notes']) == {
            'Kolaborasi', 'Kreativitas'}

    def test_regen_aktivitas_uses_fresh_principles(self, pipeline):
        """UAT FIX: regen aktivitas memakai prinsip fresh dari responsnya,
        bukan carry-over basi."""
        import app as app_module
        from module_store import ModuleStore
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai,
            mock_meeting_principles)
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        store = ModuleStore(pipeline.db_path)
        store.save('gid-prin-test', module,
                   module.get('curriculum_context') or {})
        app_module._generation_pipeline = pipeline
        try:
            fresh = mock_meeting_principles(2)
            for entry in fresh:
                for key in ('berkesadaran', 'bermakna', 'menggembirakan'):
                    entry[key] = entry[key] + ' FRESH.'

            def regen_mock(prompt, system_instruction=None, **kwargs):
                if '"activities"' in prompt:
                    resp = _mock_activities()
                    resp['meeting_principles'] = fresh
                    return resp
                return fake_router_response(
                    prompt, system_instruction, **kwargs)

            with patch_pipeline_ai(pipeline, regen_mock):
                client = app_module.app.test_client()
                resp = client.post('/api/module/gid-prin-test/regenerate',
                                   json={'base_version': 1,
                                         'target': 'aktivitas'})
            assert resp.status_code == 200, resp.get_data(as_text=True)
            import json as _json
            meetings = _json.loads(resp.data)['version']['module']['meetings']
            texts = [m.get('principles', {}).get('berkesadaran', '')
                     for m in meetings]
            assert all('FRESH.' in t for t in texts), texts
        finally:
            app_module._generation_pipeline = None

    def test_generate_does_not_mutate_input(self, pipeline):
        """Regresi konkret Batch 8.5: generate() menulis kunci turunan ke
        params; dict milik pemanggil (mis. global VALID_PARAMS) tidak
        boleh berubah — dua generate beruntun harus identik hasilnya."""
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        params = dict(MEET_PARAMS)

        def side_effect(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _mock_activities()
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(pipeline, side_effect):
            first = pipeline.generate(params)
            assert first['status'] == 'success', first['errors']
            assert params['topic'] == 'Zakat dan wakaf', params
            assert 'topic_seed' not in params
            second = pipeline.generate(params)
            assert second['status'] == 'success', second['errors']
        assert first['module']['topic'] == second['module']['topic']
        assert first['module']['topic_seed'] == 'Zakat dan wakaf'

    def test_empty_seed_skips_refinement(self, pipeline):
        params = dict(MEET_PARAMS)
        params.pop('topic')
        result = _generate(pipeline, params)
        assert result['status'] == 'success', result['errors']
        assert result['module']['topic'] == ''
        assert result['module']['topic_seed'] == ''

    def test_topic_edit_keeps_seed(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        # Jalur versions: topic dapat diedit, seed immutable.
        import app as app_module
        from module_store import ModuleStore
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        store = ModuleStore(pipeline.db_path)
        store.save('gid-topic-test', module,
                   module.get('curriculum_context') or {})
        app_module._generation_pipeline = pipeline
        try:
            client = app_module.app.test_client()
            resp = client.post('/api/module/gid-topic-test/versions',
                               json={'base_version': 1,
                                     'changes': {'topic': 'Topik edit guru'}})
        finally:
            app_module._generation_pipeline = None
        assert resp.status_code == 200, resp.get_data(as_text=True)
        import json as _json
        saved = _json.loads(resp.data)['version']['module']
        assert saved['topic'] == 'Topik edit guru'
        assert saved['topic_seed'] == 'Zakat dan wakaf'


class TestPrinciplesPerMeeting:
    """§6: prinsip per pertemuan; bukan tahap; bukan mapping."""

    def test_meetings_carry_principles(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        meetings = result['module']['meetings']
        assert len(meetings) == 2
        for mtg in meetings:
            prin = mtg.get('principles') or {}
            for key in ('berkesadaran', 'bermakna', 'menggembirakan'):
                assert isinstance(prin.get(key), str) \
                    and len(prin[key].strip()) >= 20, (mtg['index'], key)

    def test_no_mechanical_mapping(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        for act in result['module']['learning_activities']:
            assert act['experience'] in (
                'memahami', 'mengaplikasi', 'merefleksi')
            assert act['stage'] in ('Pembuka', 'Inti', 'Penutup')

    def test_docx_principles_block(self, pipeline):
        import tempfile
        import os
        from docx import Document
        from docx_renderer import render_final_rpm_docx
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        text = _docx_text(module)
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            out = render_final_rpm_docx(
                {'status': 'success', 'module': module,
                 'validation': module.get('validation_results', {})},
                path)
            doc = Document(out)
            hits = sum(
                1 for t in doc.tables for r in t.rows
                if any('Prinsip Pembelajaran Mendalam' in c.text
                       for c in r.cells))
            assert hits == 2, hits
        finally:
            os.unlink(path)
        for label in ('Berkesadaran', 'Bermakna', 'Menggembirakan'):
            assert label in text

    def test_missing_principles_fails_closed(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def no_principles(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                resp = _mock_activities()
                resp.pop('meeting_principles', None)
                return resp
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(pipeline, no_principles):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_identical_principles_rejected(self):
        from module_generation_pipeline import (
            ModuleGenerationPipeline as P)
        entry = {'meeting': 1, 'berkesadaran': 'x' * 25, 'bermakna': 'x' * 25,
                 'menggembirakan': 'x' * 25}
        with pytest.raises(ValueError):
            P._validate_meeting_principles(
                [dict(entry), dict(entry, meeting=2)], 2)


class TestAssessmentCounts:
    """§7: default, custom, invalid."""

    def test_defaults_10_10_5(self, pipeline):
        params = {k: v for k, v in MEET_PARAMS.items()}
        result = _generate(pipeline, params)
        assert result['status'] == 'success', result['errors']
        asm = result['module']['assessments']
        assert len(asm['diagnostic']) == 10
        assert len(asm['formative']) == 10
        assert len(asm['summative']['items']) == 5

    def test_custom_counts(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def side_effect(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _mock_activities()
            return fake_router_response(prompt, system_instruction, **kwargs)

        params = dict(MEET_PARAMS,
                      assessment_counts={'diagnostic': 2, 'formative': 3,
                                         'summative': 3})
        with patch_pipeline_ai(pipeline, side_effect):
            result = pipeline.generate(params)
        assert result['status'] == 'success', result['errors']
        asm = result['module']['assessments']
        assert len(asm['diagnostic']) == 2
        assert len(asm['formative']) == 3
        assert len(asm['summative']['items']) == 3

    def test_invalid_counts_rejected(self, pipeline):
        # R-42: 0 = bucket tidak digenerate (valid bila sisa bucket > 0).
        for bad in ({'diagnostic': 0, 'formative': 0, 'summative': 0},
                    {'formative': 99},
                    {'summative': 'banyak'}, {'diagnostic': -2}):
            params = dict(MEET_PARAMS, assessment_counts=bad)
            result = _generate(pipeline, params)
            assert result['status'] == 'failed', bad
            assert result['module'] is None

    def test_zero_count_skips_bucket(self, pipeline):
        # R-42: bucket 0 tidak digenerate, sisa bucket tetap ada.
        # Formatif 3 soal agar semua TP terlink (mock round-robin).
        params = dict(MEET_PARAMS,
                      assessment_counts={'diagnostic': 0, 'formative': 3,
                                         'summative': 0})
        result = _generate(pipeline, params)
        assert result['status'] == 'success', result['errors']
        asm = result['module']['assessments']
        assert 'diagnostic' not in asm
        assert len(asm['formative']) == 3
        assert 'summative' not in asm

    def test_count_mismatch_fails_closed(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def short(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _mock_activities()
            resp = fake_router_response(prompt, system_instruction, **kwargs)
            if '"diagnostic":' in prompt and isinstance(resp, dict):
                resp = dict(resp)
                resp['formative'] = (resp.get('formative') or [])[:2]
            return resp

        params = dict(MEET_PARAMS,
                      assessment_counts={'diagnostic': 2, 'formative': 3,
                                         'summative': 3})
        with patch_pipeline_ai(pipeline, short):
            result = pipeline.generate(params)
        assert result['status'] == 'failed'
        assert result['module'] is None


class TestLkpdRemoved:
    """§8: tanpa section LKPD; A–E utuh."""

    def test_no_lkpd_section_docx(self, pipeline):
        result = _generate(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'LKPD' not in text
        for section in ('A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                        'C. DESAIN PEMBELAJARAN',
                        'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN',
                        'E. ASESMEN'):
            assert section in text
