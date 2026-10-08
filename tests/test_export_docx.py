#!/usr/bin/env python3
"""Test endpoint export DOCX (GET /api/module/<id>/export-docx).

Endpoint hanya melayani FINAL VALIDATED RPM melalui renderer gate;
draft/gagal/tidak dikenal ditolak tanpa file.
"""
import copy
import io
import json
import zipfile

import pytest


@pytest.fixture
def api_pipeline(test_db):
    """Flask app wired to the real pipeline on the temp test DB
    (same pattern as test_api.py)."""
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _final_module(flask_client, api_pipeline):
    """Generate modul final-validated via pipeline mock + simpan ke store."""
    import app as app_module
    from test_phase_c_module_generation import (
        fake_router_response, mock_meeting_principles, patch_pipeline_ai)
    from module_store import ModuleStore
    context_data = {
        'education_system': 'KEMENAG',
        'institution_type': 'MTs',
        'grade': 'MTs_7',
        'subject': 'Akidah Akhlak',
        'element': 'Pemahaman Konsep',
        'phase': 'D',
        'topic': 'Taubat',
    }
    rows = [
        (1, 'Pembuka', 15, 'TP-1', 'memahami'),
        (1, 'Inti', 35, 'TP-1', 'mengaplikasi'),
        (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
        (2, 'Pembuka', 15, 'TP-2', 'memahami'),
        (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
        (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
    ]

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            return {'activities': [
                {'name': f'A{i + 1}',
                 'description': 'Deskripsi ' + f'A{i + 1}',
                 'duration': dur, 'experience': exp, 'tp_linked': tp,
                 'stage': stage, 'meeting': meeting}
                for i, (meeting, stage, dur, tp, exp) in enumerate(rows)],
                'meeting_principles': mock_meeting_principles(2)}
        return fake_router_response(prompt, system_instruction, **kwargs)

    ctx = flask_client.post('/api/context/generate', json=context_data)
    assert ctx.status_code == 200, ctx.get_data(as_text=True)
    with patch_pipeline_ai(api_pipeline, side_effect):
        resp = flask_client.post('/api/module/generate', json=dict(
            context_data,
            jumlah_pertemuan=2, jp_per_pertemuan=2))
    assert resp.status_code == 200, resp.get_data(as_text=True)
    data = json.loads(resp.data)
    module = data['module']
    store = ModuleStore(api_pipeline.db_path)
    store.save(module['generation_id'], module,
               module.get('curriculum_context') or {})
    return module


def _docx_text(payload: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(payload))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return '\n'.join(texts)


@pytest.mark.api
class TestExportDocxEndpoint:
    """GET /api/module/<generation_id>/export-docx."""

    def test_export_success_final_validated(self, flask_client, api_pipeline):
        module = _final_module(flask_client, api_pipeline)
        gid = module['generation_id']
        resp = flask_client.get(f'/api/module/{gid}/export-docx')
        assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
        assert resp.content_type == (
            'application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.document')
        disp = resp.headers.get('Content-Disposition', '')
        assert 'filename=' in disp and disp.rstrip('"').endswith('.docx')
        assert zipfile.is_zipfile(io.BytesIO(resp.data))

    def test_export_content_matches_final_module(
            self, flask_client, api_pipeline):
        module = _final_module(flask_client, api_pipeline)
        gid = module['generation_id']
        resp = flask_client.get(f'/api/module/{gid}/export-docx')
        assert resp.status_code == 200
        text = _docx_text(resp.data)
        for section in ('A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                        'C. DESAIN PEMBELAJARAN', 'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN',
                        'E. ASESMEN'):
            assert section in text
        cp = module['master_outline']['design']['cp'] or ''
        assert cp[:60] in text
        tp = module['master_outline']['design']['tp'][0]['text']
        assert tp[:60] in text

    def test_export_rejected_when_final_failed(
            self, flask_client, api_pipeline):
        from module_store import ModuleStore
        module = _final_module(flask_client, api_pipeline)
        bad = copy.deepcopy(module)
        bad['generation_id'] = 'gid-failed-001'
        bad['validation_results'] = {
            'final': {'passed': False, 'errors': ['x']}}
        store = ModuleStore(api_pipeline.db_path)
        store.save('gid-failed-001', bad,
                   bad.get('curriculum_context') or {})
        resp = flask_client.get('/api/module/gid-failed-001/export-docx')
        assert resp.status_code == 422
        assert not resp.data.lstrip().startswith(b'PK')

    def test_export_rejected_when_running(
            self, flask_client, api_pipeline):
        import app as app_module
        app_module.generation_jobs['job-running-1'] = {
            'status': 'running', 'result': {}}
        try:
            resp = flask_client.get(
                '/api/module/job-running-1/export-docx')
        finally:
            app_module.generation_jobs.pop('job-running-1', None)
        assert resp.status_code == 404

    def test_export_rejected_when_unknown(self, flask_client, api_pipeline):
        resp = flask_client.get('/api/module/no-such-id/export-docx')
        assert resp.status_code == 404

    def test_export_needs_no_raw_json(self, flask_client, api_pipeline):
        """Respons export adalah biner DOCX, bukan JSON mentah."""
        module = _final_module(flask_client, api_pipeline)
        gid = module['generation_id']
        resp = flask_client.get(f'/api/module/{gid}/export-docx')
        assert resp.status_code == 200
        assert b'"tp_list"' not in resp.data
