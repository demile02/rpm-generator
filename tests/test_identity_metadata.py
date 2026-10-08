#!/usr/bin/env python3
"""Test provenance metadata identitas (semester/tahun/satuan/model).

Tanpa input user: tidak ada default diam-diam, tidak ada baris Model
pembelajaran formal. Dengan input eksplisit: nilai terbawa utuh ke
module, preview, dan DOCX.
"""
import copy
import io
import json
import zipfile

import pytest


@pytest.fixture
def api_pipeline(test_db):
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _mocked_module(api_pipeline, flask_client, body_extra=None):
    import app as app_module  # noqa
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)
    context_data = {
        'education_system': 'KEMENAG',
        'institution_type': 'MTs',
        'grade': 'MTs_7',
        'subject': 'Akidah Akhlak',
        'element': 'Pemahaman Konsep',
        'phase': 'D',
        'topic': 'Taubat',
    }
    ctx = flask_client.post('/api/context/generate', json=context_data)
    assert ctx.status_code == 200, ctx.get_data(as_text=True)
    body = dict(context_data, **dict(body_extra or {}))
    with patch_pipeline_ai(api_pipeline, fake_router_response):
        resp = flask_client.post('/api/module/generate', json=body)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    return json.loads(resp.data)


@pytest.mark.api
class TestIdentityMetadataProvenance:
    """Tanpa input -> None/kosong (bukan angka/teks karangan)."""

    def test_no_silent_defaults(self, flask_client, api_pipeline):
        data = _mocked_module(api_pipeline, flask_client)
        ident = data['module']['module_identity']
        assert ident.get('semester') is None
        assert ident.get('year') is None
        assert ident.get('satuan_pendidikan') is None
        assert data['module'].get('learning_model') in ('', None)

    def test_explicit_identity_honored(self, flask_client, api_pipeline):
        data = _mocked_module(api_pipeline, flask_client, {
            'semester': '2', 'tahun_ajaran': '2026/2027',
            'satuan_pendidikan': 'MTs Negeri 1 Contoh'})
        ident = data['module']['module_identity']
        assert ident.get('semester') == 2
        assert ident.get('year') == '2026/2027'
        assert ident.get('satuan_pendidikan') == 'MTs Negeri 1 Contoh'

    def test_semester_normalization(self, flask_client, api_pipeline):
        import app as app_module
        assert app_module._to_semester('1') == 1
        assert app_module._to_semester(2) == 2
        assert app_module._to_semester('') is None
        assert app_module._to_semester(None) is None
        assert app_module._to_semester('3') is None
        assert app_module._to_semester('ganjil') is None

    def test_docx_no_model_row_conditional_identity(
            self, flask_client, api_pipeline):
        import os
        import tempfile
        from docx import Document
        plain = _mocked_module(api_pipeline, flask_client)
        full = _mocked_module(api_pipeline, flask_client, {
            'semester': 1, 'tahun_ajaran': '2026/2027',
            'satuan_pendidikan': 'MTs Negeri 1 Contoh'})

        def render(module_dict):
            from docx_renderer import render_final_rpm_docx
            fd, path = tempfile.mkstemp(suffix='.docx')
            os.close(fd)
            try:
                out = render_final_rpm_docx(
                    {'status': 'success', 'module': module_dict,
                     'validation': module_dict.get('validation_results', {})},
                    path)
                doc = Document(out)
                texts = [p.text for p in doc.paragraphs]
                for table in doc.tables:
                    for row in table.rows:
                        texts.extend(c.text for c in row.cells)
                return '\n'.join(texts)
            finally:
                os.unlink(path)

        text_plain = render(plain['module'])
        assert 'Model pembelajaran' not in text_plain
        # Label template statis, tetapi nilai semester tidak dikarang:
        # baris KFS hanya memuat grade/phase ramah tampil yang nyata.
        assert '7/D' in text_plain
        assert 'MTs_7' not in text_plain
        assert '7/D/' not in text_plain
        assert 'Ganjil' not in text_plain and 'Genap' not in text_plain
        assert 'Discovery Learning' not in text_plain
        text_full = render(full['module'])
        assert 'Model pembelajaran' not in text_full
        assert 'MTs Negeri 1 Contoh' in text_full
        assert '2026/2027' in text_full

    def test_preview_no_model_row(self, flask_client, api_pipeline):
        data = _mocked_module(api_pipeline, flask_client)
        # render.js difilter: baris falsy tidak tampil; pastikan kontrak
        # data untuk preview tidak membawa model formal.
        assert data['module'].get('learning_model') in ('', None)
