#!/usr/bin/env python3
"""Batch 4 §11: workflow backend — async generate → poll → result/export,
failure terminal tanpa auto-regeneration, error actionable."""
import json
import sys
import time
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

VALID_CONTEXT = {
    'education_system': 'KEMENAG',
    'institution_type': 'MTs',
    'grade': 'MTs_7',
    'subject': 'Akidah Akhlak',
    'phase': 'D',
    'element': 'Pemahaman Konsep',
    'topic': 'Taubat',
}

ROWS_D = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-1', 'mengaplikasi'),
    (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'memahami'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
]


@pytest.fixture
def api_pipeline(test_db):
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _ctx(flask_client):
    resp = flask_client.post('/api/context/generate', json=dict(VALID_CONTEXT))
    assert resp.status_code == 200, resp.get_data(as_text=True)


def _meetings_side_effect():
    from test_phase_c_module_generation import (
        fake_router_response, mock_meeting_principles)

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            return {'activities': [
                {'name': f'A{i + 1}', 'description': 'Deskripsi ' + f'A{i + 1}',
                 'duration': dur, 'experience': exp, 'tp_linked': tp,
                 'stage': stage, 'meeting': meeting}
                for i, (meeting, stage, dur, tp, exp) in enumerate(ROWS_D)],
                'meeting_principles': mock_meeting_principles(2)}
        return fake_router_response(prompt, system_instruction, **kwargs)

    return side_effect


def _poll_until(flask_client, job_id, timeout_s=120):
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        resp = flask_client.get(f'/api/generation/status/{job_id}')
        assert resp.status_code == 200
        body = json.loads(resp.data)
        if body.get('status') in ('success', 'error'):
            return body
        time.sleep(1)
    raise AssertionError('polling tidak selesai (stuck)')


class TestAsyncWorkflow:
    def test_full_flow_success(self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai
        _ctx(flask_client)
        # Patch HARUS aktif selama worker thread berjalan (poll di dalam
        # context), kalau tidak thread memakai gateway real.
        with patch_pipeline_ai(api_pipeline, _meetings_side_effect()):
            resp = flask_client.post('/api/module/generate', json={
                'async': True, **dict(VALID_CONTEXT,
                                      jumlah_pertemuan=2,
                                      jp_per_pertemuan=2)})
            assert resp.status_code == 202
            job_id = json.loads(resp.data)['job_id']
            body = _poll_until(flask_client, job_id)
        assert body['status'] == 'success', body
        module = body['module']
        assert module['subject'] == 'Akidah Akhlak'
        # Result: A–E dapat dirender dari modul final.
        from docx_renderer import render_final_rpm_docx
        import tempfile
        import os
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            out = render_final_rpm_docx(
                {'status': 'success', 'module': module,
                 'validation': body.get('validation', {})}, path)
            from docx import Document
            doc = Document(out)
            full = '\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    full += '\n' + '\n'.join(c.text for c in row.cells)
            assert 'A. IDENTITAS MODUL' in full and 'E. ASESMEN' in full
        finally:
            os.unlink(path)
        # Export endpoint melayani hasil final yang sama.
        gid = body['generation_id']
        exp = flask_client.get(f'/api/module/{gid}/export-docx')
        assert exp.status_code == 200
        assert exp.content_type.startswith(
            'application/vnd.openxmlformats')

    def test_failure_is_terminal_no_autoregen(self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai

        def always_fail(prompt, system_instruction=None, **kwargs):
            raise ValueError('mock AI down')

        _ctx(flask_client)
        with patch_pipeline_ai(api_pipeline, always_fail):
            resp = flask_client.post('/api/module/generate', json={
                'async': True, **dict(VALID_CONTEXT,
                                      jumlah_pertemuan=2,
                                      jp_per_pertemuan=2)})
            assert resp.status_code == 202
            job_id = json.loads(resp.data)['job_id']
            body = _poll_until(flask_client, job_id)
        assert body['status'] == 'error'
        assert body.get('errors'), 'pesan error wajib ada'
        # Poll ulang: tetap error yang sama, tidak ada job/generation baru.
        again = json.loads(flask_client.get(
            f'/api/generation/status/{job_id}').data)
        assert again['status'] == 'error'
        assert again.get('generation_id') == body.get('generation_id')

    def test_meetings_over_cap_actionable_400(self, flask_client, api_pipeline):
        _ctx(flask_client)
        resp = flask_client.post('/api/module/generate', json={
            'async': True, **dict(VALID_CONTEXT,
                                  jumlah_pertemuan=17,
                                  jp_per_pertemuan=2)})
        assert resp.status_code == 400
        body = json.loads(resp.data)
        assert any('16' in e.get('message', '') for e in body.get('errors', []))

    def test_jp_over_cap_actionable_400(self, flask_client, api_pipeline):
        _ctx(flask_client)
        resp = flask_client.post('/api/module/generate', json=dict(
            VALID_CONTEXT, jumlah_pertemuan=2, jp_per_pertemuan=11))
        assert resp.status_code == 400
        body = json.loads(resp.data)
        assert any('10' in e.get('message', '') for e in body.get('errors', []))

    def test_invalid_context_dependency_400(self, flask_client, api_pipeline):
        resp = flask_client.post('/api/module/generate', json=dict(
            VALID_CONTEXT, grade='MTs_99'))
        assert resp.status_code == 400
