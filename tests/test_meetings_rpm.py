#!/usr/bin/env python3
"""Regression tests: input guru, pertemuan, budget waktu, KBC di B.

Semua memakai mock AI (fakta integrasi dibuktikan via pipeline nyata +
validator nyata); E2E real terpisah. Tanpa ubah test lain.
"""
import copy

import pytest
import sys
from pathlib import Path

# Src imports (same sys.path convention as the other test modules).
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

MEET_PARAMS = {
    'education_system': 'KEMENAG',
    'institution_type': 'MTs',
    'grade': 'MTs_7',
    'subject': 'Akidah Akhlak',
    'element': 'Pemahaman Konsep',
    'topic': 'Taubat',
    'requested_tp_count': 3,
    'jumlah_pertemuan': 2,
    'jp_per_pertemuan': 2,
    'penyusun': 'Ibu Guru',
    'student_readiness': 'Sebagian murid paham wahyu.',
}

SCHOOL_PARAMS = {
    'education_system': 'KEMENDIKDASMEN',
    'institution_type': 'SD',
    'grade': 'SD_1',
    'subject': 'PENDIDIKAN PANCASILA',
    'element': 'Pancasila',
    'topic': 'Sila Pertama',
    'requested_tp_count': 3,
    'jumlah_pertemuan': 2,
    'jp_per_pertemuan': 2,
}


@pytest.fixture
def api_pipeline(test_db):
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _act(name, duration, tp, exp, stage, meeting):
    return {'name': name, 'description': 'Deskripsi ' + name,
            'duration': duration, 'experience': exp, 'tp_linked': tp,
            'stage': stage, 'meeting': meeting}


def _meetings_response(rows):
    from test_phase_c_module_generation import mock_meeting_principles
    return {'activities': [
        _act(f'A{i + 1}', dur, tp, exp, stage, meeting)
        for i, (meeting, stage, dur, tp, exp) in enumerate(rows)],
        'meeting_principles': mock_meeting_principles(2)}


# Fase D = 40 mnt/JP -> 2 pertemuan x 2 JP = 80 mnt/pertemuan.
ROWS_D = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-1', 'mengaplikasi'),
    (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'memahami'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
]

# Fase A = 35 mnt/JP -> 70 mnt/pertemuan.
ROWS_A = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-2', 'mengaplikasi'),
    (1, 'Inti', 20, 'TP-1', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'memahami'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 20, 'TP-3', 'merefleksi'),
]


def _meetings_fake(rows):
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            return _meetings_response(rows)
        return fake_router_response(prompt, system_instruction, **kwargs)

    return side_effect


@pytest.fixture
def pipeline(test_db):
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


@pytest.fixture
def api_pipeline(test_db):
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _generate(pipeline, params, rows):
    from test_phase_c_module_generation import patch_pipeline_ai
    with patch_pipeline_ai(pipeline, _meetings_fake(rows)):
        return pipeline.generate(dict(params))


def _docx_text(module_dict):
    import io
    from docx import Document
    from docx_renderer import render_final_rpm_docx
    buf = io.BytesIO()
    import tempfile
    import os
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    try:
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_dict,
             'validation': module_dict.get('validation_results', {})},
            path)
        doc = Document(out)
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.extend(c.text for c in row.cells)
        return '\n'.join(texts)
    finally:
        os.unlink(path)


class TestGuruInputs:
    """1-3: input guru diteruskan sampai output."""

    def test_penyusun_passthrough(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        assert result['module']['penyusun'] == 'Ibu Guru'

    def test_satuan_passthrough(self, pipeline):
        params = dict(MEET_PARAMS,
                      satuan_pendidikan='MTs Negeri 1 Contoh')
        result = _generate(pipeline, params, ROWS_D)
        assert result['status'] == 'success', result['errors']
        assert result['module']['module_identity'][
            'satuan_pendidikan'] == 'MTs Negeri 1 Contoh'

    def test_kesiapan_passthrough(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        assert module['student_readiness'] == 'Sebagian murid paham wahyu.'
        assert module['master_outline']['identification'][
            'learner_readiness']['summary'] == 'Sebagian murid paham wahyu.'


class TestNoSilentDefaults:
    """4-5: tanpa input pertemuan -> mode lama, tanpa angka karangan."""

    def test_no_meetings_no_fabrication(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        params = {k: v for k, v in MEET_PARAMS.items()
                  if k not in ('jumlah_pertemuan', 'jp_per_pertemuan')}
        with patch_pipeline_ai(pipeline, fake_router_response):
            result = pipeline.generate(params)
        assert result['status'] == 'success', result['errors']
        assert result['module']['meetings'] == []
        assert 'time_allocation' not in result['module']
        assert 'total_minutes' not in result['module']


class TestTimeBudget:
    """6-10: konfigurasi JP + validator budget deterministik."""

    def test_jp_config_sma_ma(self):
        from module_generation_pipeline import (
            JP_MENIT_REGULER, resolve_meetings_spec)
        assert JP_MENIT_REGULER['E'] == 45
        assert JP_MENIT_REGULER['F'] == 45
        spec, err = resolve_meetings_spec(
            {'jumlah_pertemuan': 2, 'jp_per_pertemuan': 2}, 'E', 'MA')
        assert err is None
        assert spec['minutes_per_jp'] == 45

    def test_2x2_180_sma_ma(self):
        from module_generation_pipeline import resolve_meetings_spec
        spec, err = resolve_meetings_spec(
            {'jumlah_pertemuan': 2, 'jp_per_pertemuan': 2}, 'E', 'SMA')
        assert err is None
        assert spec['total_minutes'] == 180
        assert spec['minutes_per_meeting'] == 90

    def test_partial_meetings_rejected(self, pipeline):
        bad = dict(MEET_PARAMS, jp_per_pertemuan=None)
        result = pipeline.generate(bad)
        assert result['status'] == 'failed'
        assert any('alokasi waktu' in str(e).lower()
                   for e in result['errors'])

    def test_invalid_meetings_rejected(self, pipeline):
        bad = dict(MEET_PARAMS, jumlah_pertemuan='dua')
        result = _generate(pipeline, bad, ROWS_D)
        assert result['status'] == 'failed'
        assert any('alokasi waktu' in str(e).lower()
                   for e in result['errors'])

    def test_per_meeting_budget(self):
        from module_generation_pipeline import validate_time_budget
        acts = [dict(zip(('id', 'meeting', 'stage', 'duration', 'tp_linked'),
                         (f'A{i}', m, s, d, t)))
                for i, (m, s, d, t, _e) in enumerate(ROWS_D)]
        assert validate_time_budget(acts, 2, 80, {'TP-1', 'TP-2', 'TP-3'}) == []

    def test_total_mismatch_rejected(self):
        from module_generation_pipeline import validate_time_budget
        acts = [dict(zip(('id', 'meeting', 'stage', 'duration', 'tp_linked'),
                         (f'A{i}', m, s, d, t)))
                for i, (m, s, d, t, _e) in enumerate(ROWS_D)]
        acts[0] = dict(acts[0], duration=10)
        errors = validate_time_budget(acts, 2, 80, {'TP-1', 'TP-2', 'TP-3'})
        assert errors

    def test_bad_rows_rejected(self):
        from module_generation_pipeline import validate_time_budget
        base = dict(id='A1', meeting=1, stage='Inti', duration=80,
                    tp_linked='TP-1')
        assert validate_time_budget(
            [dict(base, stage='Istirahat')], 2, 80, {'TP-1'}) != []
        assert validate_time_budget(
            [dict(base, meeting=9)], 2, 80, {'TP-1'}) != []
        assert validate_time_budget(
            [dict(base, tp_linked='TP-9')], 2, 80, {'TP-1'}) != []
        assert validate_time_budget(
            [dict(base, duration='lama')], 2, 80, {'TP-1'}) != []

    def test_pipeline_rejects_bad_budget(self, pipeline):
        rows = [list(r) for r in ROWS_D]
        rows[0] = (1, 'Pembuka', 5, 'TP-1', 'memahami')
        result = _generate(pipeline, MEET_PARAMS, rows)
        assert result['status'] == 'failed'
        assert any('Alokasi waktu' in str(e) for e in result['errors'])


class TestMeetingStructure:
    """11-14: D per pertemuan, tahap + pengalaman valid, TP valid."""

    def test_d_by_meetings(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        meetings = result['module']['meetings']
        assert len(meetings) == 2
        for mtg in meetings:
            assert set(mtg['stages'].keys()) == {
                'Pembuka', 'Inti', 'Penutup'}
            assert mtg['minutes'] == 80
            assert mtg['used_minutes'] == 80

    def test_tahap_valid(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        for act in result['module']['learning_activities']:
            assert act['stage'] in ('Pembuka', 'Inti', 'Penutup')

    def test_experience_valid(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        for act in result['module']['learning_activities']:
            assert act['experience'] in (
                'memahami', 'mengaplikasi', 'merefleksi')

    def test_tp_valid(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        tp_ids = {t['id'] for t in result['module']['learning_objectives']}
        assert tp_ids == {'TP-1', 'TP-2', 'TP-3'}
        for act in result['module']['learning_activities']:
            assert act['tp_linked'] in tp_ids


class TestKbcPosition:
    """15-18: KBC di B madrasah; tanpa section pasca-E; tanpa meta."""

    def test_kbc_in_b_madrasah(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'Kurikulum Berbasis Cinta' in text
        assert text.index('Kurikulum Berbasis Cinta') < text.index(
            'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN')

    def test_no_post_e_kbc_section(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert text.count('Kurikulum Berbasis Cinta') == 1

    def test_no_meta_note(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'layer tambahan' not in text.lower()

    def test_school_no_kbc(self, pipeline):
        result = _generate(pipeline, SCHOOL_PARAMS, ROWS_A)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'Kurikulum Berbasis Cinta' not in text
        assert 'Panca Cinta' not in text


class TestDocxMeetings:
    """19-20: A–E + export final-only."""

    def test_docx_a_to_e_meetings(self, pipeline):
        result = _generate(pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        for section in ('A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                        'C. DESAIN PEMBELAJARAN',
                        'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN',
                        'E. ASESMEN'):
            assert section in text
        assert 'Pertemuan 1' in text and 'Pertemuan 2' in text
        # Template §6: tanpa baris Total / Total keseluruhan.
        assert 'Total keseluruhan' not in text
        assert 'Total: ' not in text

    def test_export_meetings_final_only(self, flask_client, api_pipeline):
        from module_store import ModuleStore
        result = _generate(api_pipeline, MEET_PARAMS, ROWS_D)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        store = ModuleStore(api_pipeline.db_path)
        store.save(module['generation_id'], module,
                   module.get('curriculum_context') or {})
        resp = flask_client.get(
            f"/api/module/{module['generation_id']}/export-docx")
        assert resp.status_code == 200
        assert resp.content_type.startswith(
            'application/vnd.openxmlformats')
        bad = copy.deepcopy(module)
        bad['generation_id'] = 'gid-bad-meetings'
        bad['validation_results'] = {
            'final': {'passed': False, 'errors': ['x']}}
        store.save('gid-bad-meetings', bad,
                   bad.get('curriculum_context') or {})
        resp2 = flask_client.get('/api/module/gid-bad-meetings/export-docx')
        assert resp2.status_code == 422


class TestAppInputValidation:
    """API menolak alokasi sebagian/invalid dengan error jelas."""

    def _context(self, flask_client):
        ctx = flask_client.post('/api/context/generate', json={
            'education_system': 'KEMENAG', 'institution_type': 'MTs',
            'grade': 'MTs_7', 'subject': 'Akidah Akhlak',
            'element': 'Pemahaman Konsep', 'phase': 'D', 'topic': 'Taubat'})
        assert ctx.status_code == 200

    def test_partial_meetings_400(self, flask_client, api_pipeline):
        self._context(flask_client)
        resp = flask_client.post('/api/module/generate', json=dict(
            {'education_system': 'KEMENAG', 'institution_type': 'MTs',
             'grade': 'MTs_7', 'subject': 'Akidah Akhlak',
             'element': 'Pemahaman Konsep', 'phase': 'D',
             'topic': 'Taubat'},
            jumlah_pertemuan=2))
        assert resp.status_code == 400
        assert 'alokasi waktu' in resp.get_data(as_text=True).lower()

    def test_invalid_meetings_400(self, flask_client, api_pipeline):
        self._context(flask_client)
        resp = flask_client.post('/api/module/generate', json=dict(
            {'education_system': 'KEMENAG', 'institution_type': 'MTs',
             'grade': 'MTs_7', 'subject': 'Akidah Akhlak',
             'element': 'Pemahaman Konsep', 'phase': 'D',
             'topic': 'Taubat'},
            jumlah_pertemuan='banyak', jp_per_pertemuan=2))
        assert resp.status_code == 400



# ============================================================================
# BATCH 1 — JP SLB + FALLBACK DIMENSI + KBC TRACEABILITY
# ============================================================================

class TestJpSlb:
    """JP per satuan LB dari tabel struktur regulasi (bukan fase saja)."""

    def test_sdlb_fase_abc_30_menit(self):
        from module_generation_pipeline import resolve_jp_menit
        for phase in ('A', 'B', 'C'):
            minutes, err = resolve_jp_menit('SDLB', phase)
            assert err is None
            assert minutes == 30, (phase, minutes)

    def test_smplb_mtslb_fase_d_35_menit(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('SMPLB', 'D')
        assert err is None
        assert minutes == 35

    def test_smalb_malb_fase_ef_40_menit(self):
        from module_generation_pipeline import resolve_jp_menit
        for phase in ('E', 'F'):
            minutes, err = resolve_jp_menit('SMALB', phase)
            assert err is None
            assert minutes == 40, (phase, minutes)

    def test_reguler_slb_tidak_tertukar_fase_d(self):
        from module_generation_pipeline import resolve_jp_menit
        reguler, err1 = resolve_jp_menit('MTs', 'D')
        slb, err2 = resolve_jp_menit('SMPLB', 'D')
        assert err1 is None and err2 is None
        assert reguler == 40 and slb == 35 and reguler != slb

    def test_reguler_slb_tidak_tertukar_fase_e(self):
        from module_generation_pipeline import resolve_jp_menit
        reguler, _ = resolve_jp_menit('MA', 'E')
        slb, _ = resolve_jp_menit('SMALB', 'E')
        assert (reguler, slb) == (45, 40)

    def test_tklb_fondasi_fail_closed(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('TKLB', 'Fondasi')
        assert minutes is None
        assert err and 'belum tersedia' in err

    def test_institusi_tak_dikenal_fail_closed(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('XXX', 'D')
        assert minutes is None
        assert err and 'belum tersedia' in err

    def test_tidak_ada_fallback_slb_ke_reguler(self):
        from module_generation_pipeline import resolve_jp_menit
        # SMPLB tidak boleh jatuh ke 40 (reguler Fase D).
        minutes, err = resolve_jp_menit('SMPLB', 'D')
        assert (minutes, err) == (35, None)
        # SDLB tidak boleh jatuh ke 35 (reguler Fase A).
        minutes, err = resolve_jp_menit('SDLB', 'A')
        assert (minutes, err) == (30, None)

    def test_tidak_ada_fallback_phase_default(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('RA', 'A')
        assert minutes is None and err is not None

    def test_mechanism_grade_override(self):
        from module_generation_pipeline import resolve_jp_menit
        mapping = {
            'reguler_institutions': {'SD'},
            'reguler': {'A': 35},
            'slb': {'SDLB': {'default': {'A': 30},
                             'grades': {'SDLB_6': 33}}},
        }
        assert resolve_jp_menit('SDLB', 'A', 'SDLB_1', mapping) == (30, None)
        assert resolve_jp_menit('SDLB', 'A', 'SDLB_6', mapping) == (33, None)
        assert resolve_jp_menit('SD', 'A', 'SD_1', mapping) == (35, None)

    def test_spec_sdlb_memakai_30(self):
        from module_generation_pipeline import resolve_meetings_spec
        spec, err = resolve_meetings_spec(
            {'jumlah_pertemuan': 2, 'jp_per_pertemuan': 2,
             'institution_type': 'SDLB'},
            'B', 'SDLB')
        assert err is None
        assert spec['minutes_per_jp'] == 30
        assert spec['total_minutes'] == 120


class TestFallbackDimensi:
    """AI diam/invalid -> kosong jujur (A3_EMPTY), bukan full-8."""

    def _generate_dims(self, pipeline, mutator=None):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def side_effect(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _meetings_response(ROWS_D)
            resp = fake_router_response(
                prompt, system_instruction, **kwargs)
            if 'triggering_questions' in prompt and mutator is not None:
                resp = mutator(dict(resp))
            return resp

        with patch_pipeline_ai(pipeline, side_effect):
            return pipeline.generate(dict(MEET_PARAMS))

    def test_valid_dimensi_digunakan(self, pipeline):
        result = self._generate_dims(pipeline)
        assert result['status'] == 'success', result['errors']
        assert set(result['module']['profile_dimensions']) <= {
            'Kolaborasi', 'Kreativitas'}

    def test_invalid_dimensi_ditolak(self, pipeline):
        def mut(resp):
            resp['profile_dimensions'] = ['Dimensi Kejayaan']
            return resp
        result = self._generate_dims(pipeline, mut)
        assert result['status'] == 'failed'
        assert any('A3_EMPTY' in str(e) for e in result['errors'])

    def test_array_kosong_tidak_full_8(self, pipeline):
        def mut(resp):
            resp['profile_dimensions'] = []
            return resp
        result = self._generate_dims(pipeline, mut)
        assert result['status'] == 'failed'
        assert any('A3_EMPTY' in str(e) for e in result['errors'])

    def test_tanpa_dimensi_tidak_full_8(self, pipeline):
        def mut(resp):
            resp.pop('profile_dimensions', None)
            return resp
        result = self._generate_dims(pipeline, mut)
        assert result['status'] == 'failed'

    def test_kosong_tetap_jujur_dan_valid(self, pipeline):
        from types import SimpleNamespace
        ctx = SimpleNamespace(education_system='KEMENAG')
        assert pipeline._resolve_profile_dimensions(
            {'profile_dimensions': []}, ctx, {}) == []
        assert pipeline._resolve_profile_dimensions(
            {'profile_dimensions': ['Kolaborasi', 'Bukan Dimensi']},
            ctx, {}) == ['Kolaborasi']

    def test_panca_cinta_bukan_dimensi(self, pipeline):
        def mut(resp):
            resp['profile_dimensions'] = ['Cinta Ilmu', 'Kolaborasi']
            return resp
        result = self._generate_dims(pipeline, mut)
        assert result['status'] == 'success', result['errors']
        pd = result['module']['profile_dimensions']
        assert pd == ['Kolaborasi']
        assert 'Cinta Ilmu' not in pd

    def test_output_valid_tidak_berubah(self, pipeline):
        result = self._generate_dims(pipeline)
        assert result['status'] == 'success', result['errors']
        assert result['module']['profile_dimensions'] == [
            'Kolaborasi', 'Kreativitas']


class TestKbcTraceability:
    """Insersi -> aktivitas eksplisit via TP tervalidasi."""

    def _madrasah_result(self, pipeline, kbc_extra=None):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)

        def kbc_mock(prompt, system_instruction=None, **kwargs):
            if '"activities"' in prompt:
                return _meetings_response(ROWS_D)
            resp = fake_router_response(prompt, system_instruction, **kwargs)
            if 'triggering_questions' in prompt and kbc_extra is not None:
                resp = dict(resp)
                resp['kbc'] = dict(resp.get('kbc') or {})
                resp['kbc'].update(kbc_extra)
            return resp

        with patch_pipeline_ai(pipeline, kbc_mock):
            return pipeline.generate(dict(MEET_PARAMS))

    def test_madrasah_kbc_dipilih(self, pipeline):
        result = self._madrasah_result(pipeline)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        assert kbc['enabled'] is True
        assert kbc['themes']

    def test_tema_insersi_terhubung(self, pipeline):
        kbc_extra = {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'Diterapkan saat analisis sumber.',
        }
        result = self._madrasah_result(pipeline, kbc_extra)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        assert kbc['themes'] == ['Cinta Ilmu']
        assert len(kbc['insertions']) == 1
        assert kbc['insertions'][0]['tp_ids'] == ['TP-1']

    def test_insersi_activity_ids_valid(self, pipeline):
        kbc_extra = {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'x',
        }
        result = self._madrasah_result(pipeline, kbc_extra)
        assert result['status'] == 'success', result['errors']
        acts = {a['id'] for a in result['module']['learning_activities']}
        for ins in result['module']['approach_principles']['kbc']['insertions']:
            assert ins['activity_ids']
            assert set(ins['activity_ids']) <= acts

    def test_aktivitas_memuat_nilai(self, pipeline):
        kbc_extra = {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'x',
        }
        result = self._madrasah_result(pipeline, kbc_extra)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        tp_ids = {t['id'] for t in result['module']['learning_objectives']}
        for ins in kbc['insertions']:
            assert set(ins['tp_ids']) <= tp_ids

    def test_insersi_tanpa_aktivitas_tidak_aktif(self, pipeline):
        kbc_extra = {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tanpa tautan', 'tp_terkait': ['TP-99']}],
            'integration_note': 'x',
        }
        result = self._madrasah_result(pipeline, kbc_extra)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        assert len(kbc['insertions']) == 1
        assert kbc['insertions'][0]['active'] is False
        text = _docx_text(result['module'])
        assert 'Insersi tanpa tautan' not in text

    def test_sekolah_tanpa_kbc(self, pipeline):
        result = _generate(pipeline, SCHOOL_PARAMS, ROWS_A)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'Kurikulum Berbasis Cinta' not in text

    def test_kbc_bukan_dimensi_9(self, pipeline):
        result = self._madrasah_result(pipeline)
        assert result['status'] == 'success', result['errors']
        pd = result['module']['profile_dimensions']
        assert not any('Cinta' in d for d in pd)

    def test_kbc_tidak_setelah_e(self, pipeline):
        result = self._madrasah_result(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert text.count('Kurikulum Berbasis Cinta') == 1
        assert text.index('Kurikulum Berbasis Cinta') < text.index(
            'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN')

    def test_meta_tetap_hilang(self, pipeline):
        result = self._madrasah_result(pipeline)
        assert result['status'] == 'success', result['errors']
        assert 'layer tambahan' not in _docx_text(result['module']).lower()

    def test_meta_note_ditolak_fail_closed(self, pipeline):
        """Temuan audit Batch 7: integration_note berisi gema arsitektur
        ('layer tambahan') wajib ditolak — tidak dirender ke guru."""
        result = self._madrasah_result(pipeline, {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'KBC menjadi layer tambahan yang '
                                'memperdalam penghayatan.',
        })
        assert result['status'] == 'failed'
        assert result['module'] is None
        assert any('KBC_META_LANGUAGE' in str(e) for e in result['errors'])

    def test_meta_varian_lapisan_ditolak(self, pipeline):
        result = self._madrasah_result(pipeline, {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'KBC sebagai lapisan nilai tambahan.',
        })
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_meta_dalam_insersi_ditolak(self, pipeline):
        result = self._madrasah_result(pipeline, {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi sebagai lapisan tambahan.', 'tp_terkait': ['TP-1']}],
            'integration_note': 'Diterapkan saat analisis sumber.',
        })
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_catatan_konkret_lolos(self, pipeline):
        result = self._madrasah_result(pipeline, {
            'themes': ['Cinta Ilmu'],
            'insertion_material': [
                {'text': 'Insersi tabayyun sumber', 'tp_terkait': ['TP-1']}],
            'integration_note': 'Diterapkan saat diskusi kelompok '
                                'menganalisis sumber.',
        })
        assert result['status'] == 'success', result['errors']

    def test_output_lama_tetap_valid(self, pipeline):
        result = self._madrasah_result(pipeline)
        assert result['status'] == 'success', result['errors']
        kbc = result['module']['approach_principles']['kbc']
        assert kbc['enabled'] is True
        text = _docx_text(result['module'])
        assert 'Kurikulum Berbasis Cinta' in text

