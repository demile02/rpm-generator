#!/usr/bin/env python3
"""Batch 4: KBC berpasangan + kontrak Pilihan Ganda.

KBC mengikuti Templatku.docx (pasangan bernomor tema + insersi,
bukan dua blok daftar). Setiap multiple_choice wajib question +
>=4 opsi A-D + correct answer yang cocok, fail-closed bila invalid.
"""
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from docx_renderer import (  # noqa: E402
    pair_kbc_items, render_final_rpm_docx, _answer_label, _mc_options,
)
from module_generation_pipeline import (  # noqa: E402
    ModuleGenerationPipeline, normalize_mc_item,
)


@pytest.fixture
def pipeline(test_db):
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


def _docx_text(module_dict):
    import tempfile
    import os
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    try:
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_dict,
             'validation': module_dict.get('validation_results', {})},
            path)
        doc = Document(out)
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.extend(c.text for c in row.cells)
        return '\n'.join(texts)
    finally:
        os.unlink(path)


def _mc(**over):
    base = {
        'question': 'Apa itu taubat?',
        'type': 'multiple_choice',
        'options': [
            {'label': 'A', 'text': 'Jawaban A'},
            {'label': 'B', 'text': 'Jawaban B'},
            {'label': 'C', 'text': 'Jawaban C'},
            {'label': 'D', 'text': 'Jawaban D'},
        ],
        'correct_answer': 'B',
    }
    base.update(over)
    return base


class TestNormalizeMcItem:
    def test_valid_dict_options(self):
        out = normalize_mc_item(_mc(), 'diagnostic[0]')
        assert out['correct_answer'] == 'B'
        assert [o['label'] for o in out['options']] == list('ABCD')

    def test_string_options_get_labels(self):
        out = normalize_mc_item(
            _mc(options=['satu', 'dua', 'tiga', 'empat'],
                correct_answer='tiga'), 'diagnostic[0]')
        assert [o['label'] for o in out['options']] == list('ABCD')
        assert out['correct_answer'] == 'C'

    def test_answer_by_text_normalized_to_label(self):
        out = normalize_mc_item(_mc(correct_answer='Jawaban D'),
                                'diagnostic[0]')
        assert out['correct_answer'] == 'D'

    def test_missing_options_rejected(self):
        with pytest.raises(ValueError):
            normalize_mc_item(
                {'question': 'Q', 'type': 'multiple_choice'},
                'diagnostic[0]')

    def test_two_options_rejected(self):
        with pytest.raises(ValueError):
            normalize_mc_item(_mc(options=['A', 'B']), 'diagnostic[0]')

    def test_broken_label_rejected(self):
        bad = _mc()
        bad['options'][1]['label'] = 'Z'
        with pytest.raises(ValueError):
            normalize_mc_item(bad, 'diagnostic[0]')

    def test_invalid_answer_rejected(self):
        with pytest.raises(ValueError):
            normalize_mc_item(_mc(correct_answer='E'), 'diagnostic[0]')
        with pytest.raises(ValueError):
            normalize_mc_item(_mc(correct_answer=''), 'diagnostic[0]')

    def test_non_mc_passthrough(self):
        item = {'question': 'Jelaskan.', 'type': 'short_answer'}
        assert normalize_mc_item(item, 'formative[0]') is item


class TestMcGenerationContract:
    def test_invalid_mc_fails_closed(self, pipeline):
        """MC tanpa opsi -> bounded recovery -> failed jujur."""
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        from test_meetings_rpm import MEET_PARAMS

        def no_options(prompt, system_instruction=None, **kwargs):
            resp = fake_router_response(prompt, system_instruction,
                                        **kwargs)
            if 'diagnostic' in prompt:
                resp = dict(resp)
                resp['diagnostic'] = [
                    {'question': 'Q tanpa opsi', 'type': 'multiple_choice'}]
            return resp

        with patch_pipeline_ai(pipeline, no_options):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_valid_mc_stored_normalized(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        from test_meetings_rpm import MEET_PARAMS, ROWS_D, _meetings_fake
        with patch_pipeline_ai(pipeline, _meetings_fake(ROWS_D)):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'success', result['errors']
        diag = result['module']['master_outline']['assessment']['awal']
        items = diag if isinstance(diag, list) else diag.get('items', [])
        assert items
        for it in items:
            assert len(it['options']) >= 4
            assert [o['label'] for o in it['options']][:4] == list('ABCD')
            assert it['correct_answer'] in [
                o['label'] for o in it['options']]


class TestKbcPairing:
    def test_tema_field_pairs_exactly(self):
        pairs = pair_kbc_items(
            ['Cinta Ilmu', 'Cinta Lingkungan'],
            [{'text': 'Insersi lingkungan', 'tema': 'Cinta Lingkungan'},
             {'text': 'Insersi ilmu', 'tema': 'Cinta Ilmu'}])
        assert pairs == [('Cinta Ilmu', 'Insersi ilmu'),
                         ('Cinta Lingkungan', 'Insersi lingkungan')]

    def test_positional_when_balanced_without_tema(self):
        pairs = pair_kbc_items(['T1', 'T2'], ['I1', 'I2'])
        assert pairs == [('T1', 'I1'), ('T2', 'I2')]

    def test_unequal_never_forced(self):
        pairs = pair_kbc_items(['T1', 'T2', 'T3'], ['I1'])
        assert pairs == [('T1', None), ('T2', None), ('T3', None),
                         (None, 'I1')]

    def test_theme_without_insertion_kept(self):
        pairs = pair_kbc_items(['T1'], [])
        assert pairs == [('T1', None)]


def _kbc_module():
    import test_meetings_rpm as tmr
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)
    from module_generation_pipeline import ModuleGenerationPipeline
    import tempfile
    import shutil
    tmp = tempfile.mkdtemp()
    shutil.copy2('db/rpm_generator.db', f'{tmp}/t.db')
    shutil.copy2('db/curriculum_mappings.json',
                 f'{tmp}/curriculum_mappings.json')
    pipe = ModuleGenerationPipeline(db_path=f'{tmp}/t.db')
    try:
        meetings_fake = tmr._meetings_fake(tmr.ROWS_D)

        def kbc_active(prompt, system_instruction=None, **kwargs):
            resp = meetings_fake(prompt, system_instruction, **kwargs)
            if 'triggering_questions' in prompt:
                resp = dict(resp)
                kbc = dict(resp.get('kbc') or {})
                kbc['insertion_material'] = [
                    {'text': 'Insersi ilmu melalui diskusi kelompok',
                     'tema': 'Cinta Ilmu', 'tp_terkait': ['TP-1']},
                    {'text': 'Insersi lingkungan melalui karya nyata',
                     'tema': 'Cinta Lingkungan',
                     'tp_terkait': ['TP-2']},
                ]
                kbc['themes'] = ['Cinta Ilmu', 'Cinta Lingkungan']
                resp['kbc'] = kbc
            return resp

        with patch_pipeline_ai(pipe, kbc_active):
            result = pipe.generate(dict(tmr.MEET_PARAMS))
        assert result['status'] == 'success', result['errors']
        return result['module']
    finally:
        pipe.curriculum_validator.close()


class TestKbcRender:
    def test_no_sibling_sections_docx(self):
        text = _docx_text(_kbc_module())
        assert 'Tema yang relevan' not in text
        assert 'Materi insersi' not in text
        assert 'Kurikulum Berbasis Cinta & Materi Insersi' in text

    def test_numbered_pairs_docx(self):
        text = _docx_text(_kbc_module())
        assert '1. Cinta Ilmu' in text
        assert '2. Cinta Lingkungan' in text
        assert 'Insersi ilmu melalui diskusi kelompok' in text
        assert 'Insersi lingkungan melalui karya nyata' in text
        # Pasangan benar: insersi ilmu di bawah tema ilmu.
        assert (text.index('1. Cinta Ilmu')
                < text.index('Insersi ilmu melalui diskusi kelompok')
                < text.index('2. Cinta Lingkungan'))


class TestAnswerLabel:
    """Batch 4.1: Jawaban Benar diagnostik hanya label A/B/C/D."""

    def _opts(self):
        return [{'label': 'A', 'text': 'Alpha'},
                {'label': 'B', 'text': 'Beta'},
                {'label': 'C', 'text': 'Gamma'},
                {'label': 'D', 'text': 'Ilham yang diterima para ulama'}]

    def test_label_passthrough(self):
        assert _answer_label({'correct_answer': 'D'},
                             self._opts()) == 'D'

    def test_prefixed_text_resolves(self):
        assert _answer_label(
            {'correct_answer': 'D. Ilham yang diterima para ulama'},
            self._opts()) == 'D'

    def test_full_text_resolves(self):
        assert _answer_label({'correct_answer': 'Beta'},
                             self._opts()) == 'B'

    def test_unmatched_honest_fallback(self):
        assert _answer_label({'correct_answer': 'tidak ada di opsi'},
                             self._opts()) == 'tidak ada di opsi'

    def test_missing_empty(self):
        assert _answer_label({}, self._opts()) == ''
        assert _answer_label({'correct_answer': ''}, self._opts()) == ''

    def test_docx_column_label_only(self, pipeline):
        from test_meetings_rpm import MEET_PARAMS, ROWS_D, _meetings_fake
        from test_phase_c_module_generation import patch_pipeline_ai
        base = _meetings_fake(ROWS_D)

        def legacy_answers(prompt, system_instruction=None, **kwargs):
            resp = base(prompt, system_instruction, **kwargs)
            if 'diagnostic' in prompt:
                resp = dict(resp)
                first = dict(resp['diagnostic'][0])
                first['correct_answer'] = (
                    'Jawaban B taubat lengkap dengan penjelasan')
                first['options'] = [
                    {'label': 'A', 'text': 'Jawaban A taubat'},
                    {'label': 'B',
                     'text': 'Jawaban B taubat lengkap dengan penjelasan'},
                    {'label': 'C', 'text': 'Jawaban C taubat'},
                    {'label': 'D', 'text': 'Jawaban D taubat'},
                ]
                resp['diagnostic'] = [first] + resp['diagnostic'][1:]
            return resp

        with patch_pipeline_ai(pipeline, legacy_answers):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        assert 'Jawaban B taubat lengkap dengan penjelasan' in text  # opsi utuh
        lines = [ln.strip() for ln in text.splitlines()]
        jawab = [ln for ln in lines if ln in ('A', 'B', 'C', 'D')]
        assert jawab and set(jawab) <= {'A', 'B', 'C', 'D'}
        assert 'B' in jawab  # label tampil di kolom Jawaban Benar

    def test_docx_column_label_only_legacy(self):
        """Legacy full-text answer tetap tampil sebagai label (aman)."""
        assert _answer_label(
            {'correct_answer': 'Kalam Allah Swt. berbahasa Arab'},
            [{'label': 'A', 'text': 'Kalam Allah Swt. berbahasa Arab'},
             {'label': 'B', 'text': 'Disusun Nabi'},
             {'label': 'C', 'text': 'Karangan sahabat'},
             {'label': 'D', 'text': 'Mimpi ulama'}]) == 'A'


class TestMcRender:
    def test_docx_shows_options_and_answer(self, pipeline):
        from test_phase_c_module_generation import (
            fake_router_response, patch_pipeline_ai)
        from test_meetings_rpm import MEET_PARAMS, ROWS_D, _meetings_fake
        with patch_pipeline_ai(pipeline, _meetings_fake(ROWS_D)):
            result = pipeline.generate(dict(MEET_PARAMS))
        assert result['status'] == 'success', result['errors']
        text = _docx_text(result['module'])
        for label in ('A. Jawaban A taubat', 'B. Jawaban B taubat',
                      'C. Jawaban C taubat', 'D. Jawaban D taubat'):
            assert label in text, label
        assert 'Jawaban Benar' in text
        assert 'multiple_choice' not in text
