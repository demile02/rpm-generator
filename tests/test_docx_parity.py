#!/usr/bin/env python3
"""Parity Web Preview ↔ DOCX + struktur template Templatku.docx.

Web Preview (ui/render.js, Batch 1) adalah representasi final result;
DOCX (src/docx_renderer.py, Batch 2) wajib merender result yang SAMA
dengan layout template: banner merged ter-shade, header kolom polos,
kolom proporsional template, font Bookman Old Style, A4 margin 2,5 cm.
Karena Web merender data modul 1:1, membandingkan DOCX terhadap data
modul == membandingkan DOCX terhadap Web.
"""
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.shared import RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from docx_renderer import render_final_rpm_docx  # noqa: E402

BANNER_FILL = 'E2EFD9'
FONT = 'Bookman Old Style'

# Section legacy yang tidak tampil di Web — tidak boleh di DOCX.
DOCX_ONLY_MARKERS = [
    'Kemitraan pembelajaran',
    'Lingkungan pembelajaran',
    'Pemanfaatan digital',
    'Pemahaman bermakna',
    'Pertanyaan pemantik',
    'Rubrik penskoran',
    'Kisi-kisi (blueprint)',
    'Remedial',
    'Pengayaan',
    'prerequisites',
    'potential_misconceptions',
    'Provenance CP:',
    'Sumber CP:',
    'source_fragment_id',
]


@pytest.fixture
def pipeline(test_db):
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


@pytest.fixture
def module_meetings(pipeline):
    """Modul final-validated jalur meetings (mock AI + validator nyata)."""
    import test_meetings_rpm as tmr
    from test_phase_c_module_generation import patch_pipeline_ai
    with patch_pipeline_ai(pipeline, tmr._meetings_fake(tmr.ROWS_D)):
        result = pipeline.generate(dict(tmr.MEET_PARAMS))
    assert result['status'] == 'success', result['errors']
    assert result['validation']['final']['passed'] is True
    return result['module']


def _render(module, tmp_path, name='rpm.docx'):
    out = render_final_rpm_docx(
        {'status': 'success', 'module': module,
         'validation': module.get('validation_results', {})},
        tmp_path / name)
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return doc, '\n'.join(texts)


def _cell_fill(cell):
    fills = cell._tc.xpath('.//w:shd/@w:fill')
    return fills[0].upper() if fills else None


def _tabel_berbanner(doc, banner):
    """Cari tabel section berdasarkan teks banner (tahan sisipan tabel)."""
    for table in doc.tables:
        if table.rows and table.rows[0].cells[0].text == banner:
            return table
    raise AssertionError('banner tidak ketemu: ' + banner)


def _tabel_e(doc):
    return _tabel_berbanner(doc, 'E. ASESMEN')
    fills = cell._tc.xpath('.//w:shd/@w:fill')
    return fills[0].upper() if fills else None


def _row_texts(table):
    return [[c.text for c in row.cells] for row in table.rows]


# ------------------------------------------------------------------
# Content parity (Web == DOCX).
# ------------------------------------------------------------------

class TestContentParity:
    def test_header_topic_match_web(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        assert 'PERENCANAAN PEMBELAJARAN MENDALAM' in text
        topik = (module_meetings.get('topic') or '').strip()
        assert topik and topik in text

    def test_section_order_a_to_e(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        sections = ['A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                    'C. DESAIN PEMBELAJARAN',
                    'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN', 'E. ASESMEN']
        for section in sections:
            assert section in text, section
        positions = [text.index(s) for s in sections]
        assert positions == sorted(positions)

    def test_identity_fields_match_web(self, module_meetings, tmp_path):
        # Template §2: label-value, tanpa Versi Kurikulum, kode ramah.
        _, text = _render(module_meetings, tmp_path)
        for label in ['Nama Penyusun', 'Mata Pelajaran',
                      'Kelas / Fase /Semester', 'Alokasi Waktu']:
            assert label in text, label
        assert 'Versi Kurikulum' not in text
        assert module_meetings['penyusun'] in text
        assert module_meetings['subject'] in text
        assert 'MTs_7' not in text  # kode mentah tidak tampil

    def test_identity_friendly_formats(self, tmp_path):
        """§2: 'MA_10/E/1' tampil '10/E/Ganjil'."""
        from docx import Document as _Document
        from docx_renderer import _render_identitas as _identitas
        module = {
            'module_identity': {'satuan_pendidikan': 'MAN X',
                                'semester': 1, 'year': '2026/2027'},
            'curriculum_context': {'education_system': 'KEMENAG'},
            'subject': 'Fikih', 'grade': 'MA_10', 'phase': 'E',
            'penyusun': 'Guru', 'unit': {}, 'meetings': [],
            'curriculum_version': 'V',
        }
        doc = _Document()
        _identitas(doc, {}, module)
        flat = '\n'.join(c.text for t in doc.tables
                         for r in t.rows for c in r.cells)
        assert '10/E/Ganjil' in flat
        assert 'MA_10' not in flat
        assert 'Tahun Pelajaran\t: 2026/2027' in flat
        assert 'Versi Kurikulum' not in flat

    def test_no_app_metadata_in_body(self, module_meetings, tmp_path):
        """§3: status aplikasi/ID/metadata internal di luar badan dokumen."""
        _, text = _render(module_meetings, tmp_path)
        for marker in ['Final validation', 'ID: MOD-', 'generation_id',
                       'prompt_hash', 'validation_results']:
            assert marker not in text, marker

    def test_topic_once(self, module_meetings, tmp_path):
        """§3: topik tepat satu kali (bawah judul, bukan section C)."""
        _, text = _render(module_meetings, tmp_path)
        topik = (module_meetings.get('topic') or '').strip()
        assert topik and text.count(topik) == 1

    def test_no_totals(self, module_meetings, tmp_path):
        """§6: tanpa baris Total/Total keseluruhan."""
        _, text = _render(module_meetings, tmp_path)
        assert 'Total keseluruhan' not in text
        assert 'Total: ' not in text

    def test_identity_school_madrasah_labels(self, tmp_path):
        """Label satuan mengikuti education_system (data, bukan default)."""
        from docx import Document as _Document
        from docx_renderer import _render_identitas as _identitas
        base = {'module_identity': {'satuan_pendidikan': 'S'},
                'subject': 'X', 'grade': 'G', 'phase': 'P',
                'penyusun': 'Y', 'unit': {}, 'curriculum_version': 'V'}
        madrasah = dict(base, curriculum_context={
            'education_system': 'KEMENAG'})
        doc = _Document()
        _identitas(doc, {}, madrasah)
        flat = '\n'.join(c.text for t in doc.tables
                         for r in t.rows for c in r.cells)
        assert 'Nama Madrasah' in flat
        sekolah = dict(base, curriculum_context={
            'education_system': 'KEMENDIKDASMEN'})
        doc2 = _Document()
        _identitas(doc2, {}, sekolah)
        flat2 = '\n'.join(c.text for t in doc2.tables
                          for r in t.rows for c in r.cells)
        assert 'Nama Sekolah' in flat2

    def test_no_docx_only_sections(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        for marker in DOCX_ONLY_MARKERS:
            assert marker not in text, marker

    def test_kktp_numbered(self, module_meetings, tmp_path):
        """§5: KKTP numbering, bukan bullet."""
        _, text = _render(module_meetings, tmp_path)
        assert '1. KKTP-1' in text

    def test_dims_numbered(self, module_meetings, tmp_path):
        """§4: dimensi numbering + uraian (bukan bullet bertingkat)."""
        _, text = _render(module_meetings, tmp_path)
        dims = module_meetings.get('profile_dimensions') or []
        assert dims
        for number, dim in enumerate(dims, start=1):
            assert f"{number}. {dim}" in text, dim

    def test_principles_numbered_with_template_labels(
            self, module_meetings, tmp_path):
        """§7: prinsip numbering + label Inggris template."""
        _, text = _render(module_meetings, tmp_path)
        for label in ['1. Berkesadaran (Mindful Learning)',
                      '2. Bermakna (Meaningful Learning)',
                      '3. Menggembirakan (Joyful Learning)']:
            assert label in text, label

    def test_diagnostic_answer_column(self, module_meetings, tmp_path):
        """§8: diagnostik memakai Jawaban Benar dari final result."""
        doc, _ = _render(module_meetings, tmp_path)
        e_table = _tabel_e(doc)
        header_idx = next(
            i for i, r in enumerate(e_table.rows)
            if [c.text for c in r.cells][:2] ==
            ['Instrumen', 'Tipe'])
        assert 'Jawaban Benar' in e_table.rows[header_idx].cells[2].text
        assert 'multiple_choice' not in '\n'.join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)

    def test_no_raw_enums(self, module_meetings, tmp_path):
        """§8: tipe tampil Bahasa Indonesia."""
        _, text = _render(module_meetings, tmp_path)
        for raw in ('multiple_choice', 'short_answer'):
            assert raw not in text, raw
        assert 'Pilihan Ganda' in text and 'Isian' in text

    def test_no_global_principles_summary(self, module_meetings, tmp_path):
        """Ringkasan 'Prinsip: ...' tidak ada di Web — tidak boleh di DOCX."""
        doc, text = _render(module_meetings, tmp_path)
        assert 'Prinsip: ' not in text
        meetings = module_meetings['meetings']
        hits = sum(
            1 for t in doc.tables for r in t.rows
            if any('Prinsip Pembelajaran Mendalam' in c.text
                   for c in r.cells))
        assert hits == len(meetings), hits

    def test_no_skor_column(self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        headers = [cell.text for table in doc.tables
                   for cell in table.rows[0].cells]
        assert 'Skor' not in headers

    def test_meeting_and_activity_counts(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        meetings = module_meetings['meetings']
        for mtg in meetings:
            assert f"Pertemuan {mtg['index']}" in text
        for mtg in meetings:
            for stage in ('Pembuka', 'Inti', 'Penutup'):
                for act in (mtg.get('stages') or {}).get(stage) or []:
                    assert (act['name'] or '')[:30] in text, act['name']

    def test_assessment_counts_match(self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        assessment = module_meetings['master_outline']['assessment']
        d_table = _tabel_e(doc)
        groups = ['Asesmen Diagnostik', 'Asesmen Formatif', 'Asesmen Sumatif']
        expected_headers = [['Instrumen', 'Tipe', 'Jawaban Benar'],
                            ['Instrumen', 'Tipe', 'TP', 'KKTP'],
                            ['Instrumen', 'Tipe', 'TP', 'KKTP']]
        for table_group, key, want in zip(
                _split_assessment(d_table), ('awal', 'proses', 'akhir'),
                expected_headers):
            header, rows = table_group
            assert header == want, (key, header)
            items = assessment.get(key) or []
            if isinstance(items, dict):
                items = items.get('items', [])
            assert len(rows) == len(items), key
        flats = [c.text for t in [d_table] for r in t.rows for c in r.cells]
        for group in groups:
            assert group in flats, group

    def test_text_parity_cp_tp_kktp_topik(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        flat = ' '.join(text.split())
        design = module_meetings['master_outline']['design']
        assert ' '.join((design['cp'] or '').split())[:60] in flat
        for t in design['tp']:
            assert (t['text'] or '')[:40] in text
        for entry in design.get('kktp') or []:
            assert entry['id'] in text
            for crit in entry.get('criteria') or []:
                assert (crit or '')[:40] in text
        assert (design.get('topic_context') or '')[:40] in text
        for practice in design.get('pedagogical_practices') or []:
            assert (practice or '')[:30] in text

    def test_principles_text_parity(self, module_meetings, tmp_path):
        _, text = _render(module_meetings, tmp_path)
        for mtg in module_meetings['meetings']:
            for key in ('berkesadaran', 'bermakna', 'menggembirakan'):
                value = ((mtg.get('principles') or {}).get(key) or '').strip()
                if value:
                    assert value[:40] in text, (mtg['index'], key)

    def test_fallback_texts_match_web(self, tmp_path):
        """Teks fallback DOCX identik dengan Web (ui/render.js)."""
        from docx import Document as _Document
        from docx_renderer import _render_desain as _desain
        doc = _Document()
        _desain(doc, {'cp': '', 'cp_provenance': {}})
        assert 'CP tidak tersedia.' in '\n'.join(
            c.text for t in doc.tables for r in t.rows for c in r.cells)

    def test_renderer_does_not_generate(self, module_meetings, tmp_path):
        src = (Path(__file__).resolve().parents[1] / 'src'
               / 'docx_renderer.py').read_text(encoding='utf-8')
        assert 'generate_json' not in src and 'ai_client' not in src
        assert 'NineRouter' not in src


def _split_assessment(table):
    """Pecah tabel E menjadi [(header, [rows])] per grup asesmen.

    Diagnostik: header 3 kolom (Jawaban Benar merged); Formatif dan
    Sumatif: header 4 kolom.
    """
    groups = []
    current = None
    for row in table.rows[1:]:
        texts = [c.text for c in row.cells]
        if texts[0].startswith('Asesmen ') and len(set(texts)) == 1:
            current = None
            continue
        # Sel merged terbaca ganda — padatkan duplikat berurutan.
        heads = []
        for t in texts:
            if t and (not heads or heads[-1] != t):
                heads.append(t)
        if heads == ['Instrumen', 'Tipe', 'Jawaban Benar']:
            current = (heads, [])
            groups.append(current)
        elif heads == ['Instrumen', 'Tipe', 'TP', 'KKTP']:
            current = (heads, [])
            groups.append(current)
        elif current is not None:
            current[1].append(texts)
    return groups


# ------------------------------------------------------------------
# Struktur template Templatku.docx.
# ------------------------------------------------------------------

class TestTemplateStructure:
    def test_page_a4_margins(self, module_meetings, tmp_path):
        from docx.shared import Cm
        doc, _ = _render(module_meetings, tmp_path)
        section = doc.sections[0]
        assert abs(section.page_width.pt - Cm(21.0).pt) < 1.0
        assert abs(section.page_height.pt - Cm(29.7).pt) < 1.0
        for margin in (section.top_margin, section.bottom_margin,
                       section.left_margin, section.right_margin):
            assert abs(margin.pt - Cm(2.5).pt) < 1.0

    def test_body_font_bookman(self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        assert doc.styles['Normal'].font.name == FONT
        fonts = {r.font.name for t in doc.tables for r in t.rows
                 for c in r.cells for p in c.paragraphs for r in p.runs}
        fonts.discard(None)
        assert fonts == {FONT}, fonts

    def test_section_banners_shaded_merged(self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        banners = ['A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                   'C. DESAIN PEMBELAJARAN',
                   'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN', 'E. ASESMEN']
        for table, banner in [( _tabel_berbanner(doc, b), b)
                               for b in banners]:
            first = [c.text for c in table.rows[0].cells]
            assert first[0] == banner, first
            if len(table.columns) > 1:
                spans = table.rows[0]._tr.xpath('.//w:gridSpan/@w:val')
                assert spans, banner  # merged seperti template
            assert _cell_fill(table.rows[0].cells[0]) == BANNER_FILL

    def test_column_headers_plain_bold(self, module_meetings, tmp_path):
        """Header kolom template: bold tanpa shade (bukan hijau)."""
        doc, _ = _render(module_meetings, tmp_path)
        tp_table = doc.tables[2]
        header = [c.text for c in tp_table.rows[3].cells]
        assert header == ['Nomor', 'Tujuan Pembelajaran', 'Level'], header
        for cell in tp_table.rows[3].cells:
            assert _cell_fill(cell) is None
            assert any(r.bold for r in cell.paragraphs[0].runs)

    def test_activity_proportions_match_template(
            self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        d_table = doc.tables[3]
        header_idx = next(
            i for i, r in enumerate(d_table.rows)
            if [c.text for c in r.cells] ==
            ['Aktivitas', 'Deskripsi', 'Durasi', 'TP', 'Pengalaman'])
        widths = [c.width.pt for c in d_table.rows[header_idx].cells]
        total = sum(widths)
        expected = [0.187, 0.422, 0.109, 0.094, 0.188]
        for got, want in zip(widths, [total * e for e in expected]):
            assert abs(got - want) < total * 0.02, widths
        assert widths[1] == max(widths)

    def test_assessment_proportions_match_template(
            self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        d_table = _tabel_e(doc)
        header_idx = next(
            i for i, r in enumerate(d_table.rows)
            if [c.text for c in r.cells] ==
            ['Instrumen', 'Tipe', 'TP', 'KKTP'])
        widths = [c.width.pt for c in d_table.rows[header_idx].cells]
        total = sum(widths)
        for got, want in zip(widths,
                             [total * e for e in
                              (0.6561, 0.1251, 0.0937, 0.1251)]):
            assert abs(got - want) < total * 0.02, widths

    def test_tables_fixed_layout_with_widths(self, module_meetings, tmp_path):
        doc, _ = _render(module_meetings, tmp_path)
        assert len(doc.tables) >= 5  # A, B, C, D, E (+ Glosarium bila ada)
        for table in doc.tables:
            assert table.autofit is False
            layouts = table._tbl.xpath('.//w:tblLayout/@w:type')
            assert layouts and layouts[0] == 'fixed'
            for row in table.rows:
                for cell in row.cells:
                    assert cell.width is not None and cell.width.pt > 0

    def test_tables_fit_usable_width(self, module_meetings, tmp_path):
        from docx.shared import Cm
        usable_pt = Cm(21.0 - 2 * 2.5).pt
        doc, _ = _render(module_meetings, tmp_path)
        for table in doc.tables:
            grid = [int(w) / 20.0 for w in
                    table._tbl.xpath('.//w:tblGrid//w:gridCol/@w:w')]
            total = sum(grid)
            assert total <= usable_pt + 1.0, total
            assert total >= usable_pt * 0.90, total

    def test_meeting_structure_principles_no_totals(
            self, module_meetings, tmp_path):
        """§6: Pertemuan → tahap (tanpa baris Total); Prinsip menutup
        tiap pertemuan di dalam tabel D (setelah Penutup)."""
        import re
        doc, _ = _render(module_meetings, tmp_path)
        d_table = doc.tables[3]
        seq = [' | '.join(c.text for c in r.cells) for r in d_table.rows]
        heads = [i for i, t in enumerate(seq)
                 if re.match(r'Pertemuan \d+:', t)]
        assert len(heads) == len(module_meetings['meetings'])
        for n, start in enumerate(heads):
            end = heads[n + 1] if n + 1 < len(heads) else len(seq)
            segment = seq[start:end]
            assert not any('Total:' in t or 'Total keseluruhan' in t
                           for t in segment)
            # Prinsip ada di akhir segmen tiap pertemuan, setelah Penutup.
            prinsip_idx = [i for i, t in enumerate(segment)
                           if 'Prinsip Pembelajaran Mendalam' in t]
            assert len(prinsip_idx) == 1, segment[-3:]
            penutup_idx = [i for i, t in enumerate(segment)
                           if t.startswith('Penutup')]
            assert penutup_idx, segment[-3:]
            assert prinsip_idx[0] > penutup_idx[-1], segment[-3:]

    def test_short_code_cells_centered(self, module_meetings, tmp_path):
        """Kolom kode pendek (Tipe/Jawaban/TP/KKTP/Nomor/Level/Durasi)
        rata tengah ala Templatku; kolom teks tetap rata kiri."""
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        doc, _ = _render(module_meetings, tmp_path)
        e_table = _tabel_e(doc)
        for row in e_table.rows:
            texts = [c.text for c in row.cells]
            if texts[:2] == ['Instrumen', 'Tipe']:
                continue
            if len(texts) >= 4 and texts[0] and not texts[0].startswith(
                    'Asesmen ') and texts[0] != 'E. ASESMEN':
                for cell in row.cells[1:]:
                    if cell.text.strip():
                        assert cell.paragraphs[0].alignment == \
                            WD_PARAGRAPH_ALIGNMENT.CENTER, cell.text[:30]

    def test_diagnostic_cells_middle_vertical(self, module_meetings,
                                              tmp_path):
        """Sel Tipe/Jawaban diagnostik tengah vertikal (Templatku)."""
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        doc, _ = _render(module_meetings, tmp_path)
        e_table = _tabel_e(doc)
        diag = False
        for row in e_table.rows:
            texts = [c.text for c in row.cells]
            if texts[0] == 'Asesmen Diagnostik':
                diag = True
                continue
            if texts[0].startswith('Asesmen ') and len(set(texts)) == 1:
                break
            if diag and texts[0] and texts[0] != 'Instrumen':
                assert row.cells[1].vertical_alignment == \
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                assert row.cells[2].vertical_alignment == \
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def test_docx_opens_without_corruption(
            self, module_meetings, tmp_path):
        import zipfile
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_meetings,
             'validation': module_meetings.get('validation_results', {})},
            tmp_path / 'rpm.docx')
        assert zipfile.is_zipfile(out)
        Document(out)  # terbuka ulang tanpa error
