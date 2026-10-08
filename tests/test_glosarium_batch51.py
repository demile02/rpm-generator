#!/usr/bin/env python3
"""Batch 5.1: glosarium — create/edit/delete, validation, versioning,
partial regen, DOCX. AI di-mock (real dibuktikan E2E terpisah)."""
import copy
import json
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

VALID_CONTEXT = {
    'education_system': 'KEMENAG',
    'institution_type': 'MTs',
    'grade': 'MTs_7',
    'subject': 'Akidah Akhlak',
    'phase': 'D',
    'element': 'Pemahaman Konsep',
    'topic': 'Taubat',
}

ROWS_D = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-1', 'mengaplikasi'),
    (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'memahami'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
]


@pytest.fixture
def api_pipeline(test_db):
    import app as app_module
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    app_module._generation_pipeline = pipe
    yield pipe
    app_module._generation_pipeline = None
    pipe.curriculum_validator.close()


def _act(name, duration, tp, exp, stage, meeting):
    return {'name': name, 'description': 'Deskripsi ' + name,
            'duration': duration, 'experience': exp, 'tp_linked': tp,
            'stage': stage, 'meeting': meeting}


def _seed(flask_client, api_pipeline):
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)
    from module_store import ModuleStore
    assert flask_client.post(
        '/api/context/generate', json=dict(VALID_CONTEXT)).status_code == 200

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            from test_phase_c_module_generation import (
                mock_meeting_principles)
            return {'activities': [
                _act(f'A{i + 1}', dur, tp, exp, stage, meeting)
                for i, (meeting, stage, dur, tp, exp) in enumerate(ROWS_D)],
                'meeting_principles': mock_meeting_principles(2)}
        return fake_router_response(prompt, system_instruction, **kwargs)

    with patch_pipeline_ai(api_pipeline, side_effect):
        resp = flask_client.post('/api/module/generate', json={
            'education_system': 'KEMENAG', 'institution_type': 'MTs',
            'grade': 'MTs_7', 'subject': 'Akidah Akhlak',
            'element': 'Pemahaman Konsep', 'phase': 'D',
            'topic': 'Taubat',
            'jumlah_pertemuan': 2, 'jp_per_pertemuan': 2})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    body = json.loads(resp.data)
    assert body['module'].get('glosarium', []) == []
    ModuleStore(api_pipeline.db_path).save(
        body['generation_id'], body['module'],
        body['module'].get('curriculum_context') or {})
    return body['generation_id'], body['module']


def _save(flask_client, gid, base, changes):
    return flask_client.post(f'/api/module/{gid}/versions',
                             json={'base_version': base, 'changes': changes})


GLOS = [
    {'istilah': 'Taubat nasuha', 'definisi': 'Kembali kepada Allah dengan '
     'menyesali dosa dan bertekad tidak mengulanginya.'},
    {'istilah': 'Istigfar', 'definisi': 'Memohon ampun kepada Allah atas '
     'dosa yang diperbuat.'},
]


class TestGlosariumEditing:
    def test_create_glossary_v2(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        resp = _save(flask_client, gid, 1, {'glosarium': GLOS})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = json.loads(resp.data)['version']
        assert body['version_no'] == 2
        assert body['changed_sections'] == ['glosarium']
        assert body['module']['glosarium'] == GLOS
        # Section lain identik dengan v1.
        v1 = json.loads(flask_client.get(
            f'/api/module/{gid}/versions/1').data)['version']['module']
        for key in ('learning_objectives', 'learning_activities',
                    'assessments', 'success_criteria'):
            assert body['module'][key] == v1[key]

    def test_edit_and_delete_entries(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        assert _save(flask_client, gid, 1, {'glosarium': GLOS}).status_code == 200
        edited = [dict(GLOS[0], definisi='Definisi baru oleh guru.')]
        resp = _save(flask_client, gid, 2, {'glosarium': edited})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = json.loads(resp.data)['version']
        assert body['module']['glosarium'] == edited
        assert body['parent_version_no'] == 2

    def test_empty_istilah_rejected_400(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        resp = _save(flask_client, gid, 1, {
            'glosarium': [{'istilah': '  ', 'definisi': 'x'}]})
        assert resp.status_code == 400

    def test_empty_definisi_rejected_400(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        resp = _save(flask_client, gid, 1, {
            'glosarium': [{'istilah': 'Taubat', 'definisi': ''}]})
        assert resp.status_code == 400

    def test_duplicate_rejected_400(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        dup = [dict(GLOS[0]), dict(GLOS[0], definisi='Definisi lain.')]
        resp = _save(flask_client, gid, 1, {'glosarium': dup})
        assert resp.status_code == 400
        assert 'duplikat' in resp.get_data(as_text=True).lower()

    def test_duplicate_case_insensitive_rejected(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        dup = [dict(GLOS[0]),
               {'istilah': 'TAUBAT NASUHA', 'definisi': 'Definisi lain.'}]
        assert _save(flask_client, gid, 1,
                     {'glosarium': dup}).status_code == 400

    def test_clear_all_allowed(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        assert _save(flask_client, gid, 1, {'glosarium': GLOS}).status_code == 200
        resp = _save(flask_client, gid, 2, {'glosarium': []})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        assert json.loads(resp.data)['version']['module']['glosarium'] == []

    def test_history_and_restore(self, flask_client, api_pipeline):
        from module_store import VersionStore
        gid, _module = _seed(flask_client, api_pipeline)
        assert _save(flask_client, gid, 1, {'glosarium': GLOS}).status_code == 200
        resp = flask_client.post(f'/api/module/{gid}/versions/1/restore')
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = json.loads(resp.data)['version']
        assert body['module'].get('glosarium', []) == []
        hist = VersionStore(api_pipeline.db_path).history(gid)
        assert [h['version_no'] for h in hist] == [1, 2, 3]


def _glosarium_mock():
    from test_phase_c_module_generation import fake_router_response

    def side_effect(prompt, system_instruction=None, **kwargs):
        if 'istilah' in prompt and 'glosarium' in prompt:
            return {'glosarium': [
                {'istilah': 'Taubat', 'definisi': 'Kembali kepada Allah '
                 'dengan penyesalan yang tulus.'},
                {'istilah': 'Muhasabah', 'definisi': 'Introspeksi diri '
                 'atas perbuatan yang telah dilakukan.'}]}
        return fake_router_response(prompt, system_instruction, **kwargs)

    return side_effect


class TestGlosariumRegen:
    def test_regen_replaces_only_glossary(self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai
        gid, module = _seed(flask_client, api_pipeline)
        snapshot = {k: copy.deepcopy(module[k]) for k in
                    ('learning_objectives', 'learning_activities',
                     'assessments', 'success_criteria', 'reflection', 'lkpd',
                     'essential_understanding')}
        with patch_pipeline_ai(api_pipeline, _glosarium_mock()):
            resp = flask_client.post(f'/api/module/{gid}/regenerate',
                                     json={'base_version': 1,
                                           'target': 'glosarium'})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = json.loads(resp.data)['version']
        assert body['changed_sections'] == ['glosarium']
        assert body['module']['glosarium'] == [
            {'istilah': 'Taubat', 'definisi': 'Kembali kepada Allah dengan '
             'penyesalan yang tulus.'},
            {'istilah': 'Muhasabah', 'definisi': 'Introspeksi diri atas '
             'perbuatan yang telah dilakukan.'}]
        for key, value in snapshot.items():
            assert body['module'][key] == value, key

    def test_regen_uses_current_version_not_baseline(
            self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai
        gid, module = _seed(flask_client, api_pipeline)
        tp = [dict(t) for t in module['learning_objectives']]
        tp[0]['text'] = 'Menerapkan taubat hasil edit guru dalam kehidupan'
        assert _save(flask_client, gid, 1, {
            'tp': [{'id': t['id'], 'text': t['text']} for t in tp]
        }).status_code == 200
        with patch_pipeline_ai(api_pipeline, _glosarium_mock()):
            resp = flask_client.post(f'/api/module/{gid}/regenerate',
                                     json={'base_version': 2,
                                           'target': 'glosarium'})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = json.loads(resp.data)['version']
        assert body['module']['learning_objectives'][0]['text'].startswith(
            'Menerapkan taubat hasil edit guru')

    def test_regen_invalid_no_new_version(self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai
        from module_store import VersionStore
        gid, _module = _seed(flask_client, api_pipeline)

        def garbage(prompt, system_instruction=None, **kwargs):
            if 'istilah' in prompt and 'glosarium' in prompt:
                return {'glosarium': 'bukan-list'}
            from test_phase_c_module_generation import fake_router_response
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(api_pipeline, garbage):
            resp = flask_client.post(f'/api/module/{gid}/regenerate',
                                     json={'base_version': 1,
                                           'target': 'glosarium'})
        assert resp.status_code == 422
        assert [h['version_no'] for h in
                VersionStore(api_pipeline.db_path).history(gid)] == [1]

    def test_regen_duplicate_no_new_version(self, flask_client, api_pipeline):
        from test_phase_c_module_generation import patch_pipeline_ai
        from module_store import VersionStore
        gid, _module = _seed(flask_client, api_pipeline)

        def dupe(prompt, system_instruction=None, **kwargs):
            if 'istilah' in prompt and 'glosarium' in prompt:
                return {'glosarium': [
                    {'istilah': 'Taubat', 'definisi': 'Definisi satu.'},
                    {'istilah': 'taubat', 'definisi': 'Definisi dua.'}]}
            from test_phase_c_module_generation import fake_router_response
            return fake_router_response(prompt, system_instruction, **kwargs)

        with patch_pipeline_ai(api_pipeline, dupe):
            resp = flask_client.post(f'/api/module/{gid}/regenerate',
                                     json={'base_version': 1,
                                           'target': 'glosarium'})
        assert resp.status_code == 422
        assert [h['version_no'] for h in
                VersionStore(api_pipeline.db_path).history(gid)] == [1]

    def test_regen_unknown_target_rejected(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        assert flask_client.post(
            f'/api/module/{gid}/regenerate',
            json={'base_version': 1, 'target': 'judul'}).status_code == 400


class TestGlosariumDocx:
    def _docx_text(self, module_dict):
        import tempfile
        import os
        from docx import Document
        from docx_renderer import render_final_rpm_docx
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            out = render_final_rpm_docx(
                {'status': 'success', 'module': module_dict,
                 'validation': module_dict.get('validation_results', {})},
                path)
            doc = Document(out)
            texts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    texts.extend(c.text for c in row.cells)
            return '\n'.join(texts)
        finally:
            os.unlink(path)

    def test_docx_contains_glossary(self, flask_client, api_pipeline):
        gid, _module = _seed(flask_client, api_pipeline)
        assert _save(flask_client, gid, 1, {'glosarium': GLOS}).status_code == 200
        text = self._docx_text(
            json.loads(flask_client.get(
                f'/api/module/{gid}/versions/2').data)['version']['module'])
        assert 'Glosarium' in text
        assert 'Taubat nasuha' in text
        assert 'Kembali kepada Allah' in text

    def test_docx_without_glossary_has_no_section(
            self, flask_client, api_pipeline):
        gid, module = _seed(flask_client, api_pipeline)
        assert 'Glosarium' not in self._docx_text(module)
