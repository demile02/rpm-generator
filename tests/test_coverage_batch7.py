#!/usr/bin/env python3
"""Batch 7 §9: coverage matrix — resolution, routing, JP, KBC, SLB,
invalid/unsupported combos, template A–E per institution.

AI di-mock (real dibuktikan E2E matrix terpisah); yang diuji adalah
otoritas kurikulum + determinisme pipeline, bukan kualitas AI.
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

# (system, institution, grade, subject, element, phase, jp, kbc?)
MATRIX = [
    ('KEMENAG', 'MI', 'MI_4', 'Akidah Akhlak', 'Pemahaman Konsep',
     'B', 35, True),
    ('KEMENDIKDASMEN', 'SMP', 'SMP_7', 'MATEMATIKA', 'Aljabar',
     'D', 40, False),
    ('KEMENAG', 'MA', 'MA_10', 'Fikih', 'Pemahaman Konsep',
     'E', 45, True),
    ('KEMENDIKDASMEN', 'SMK', 'SMK_10', 'BAHASA INDONESIA', 'Menulis',
     'E', 45, False),
    ('KEMENAG', 'MAK', 'MAK_10', 'Akidah Akhlak', 'Pemahaman Konsep',
     'E', 45, True),
    ('KEMENDIKDASMEN', 'SDLB', 'SDLB_3',
     'PENDIDIKAN KHUSUS BAHASA INDONESIA', 'Membaca dan Memirsa',
     'B', 30, False),
    ('KEMENDIKDASMEN', 'SMALB', 'SMALB_10',
     'PENDIDIKAN KHUSUS BAHASA INDONESIA', 'Membaca dan Memirsa',
     'E', 40, False),
    # Baseline yang sudah teruji (regresi matrix).
    ('KEMENAG', 'MTs', 'MTs_7', 'Akidah Akhlak', 'Pemahaman Konsep',
     'D', 40, True),
    ('KEMENDIKDASMEN', 'SMA', 'SMA_10', 'MATEMATIKA', 'Aljabar dan Fungsi',
     'E', 45, False),
    ('KEMENDIKDASMEN', 'SD', 'SD_5', 'BAHASA INDONESIA', 'Membaca dan Memirsa',
     'C', 35, False),
    ('KEMENDIKDASMEN', 'SMPLB', 'SMPLB_7',
     'PENDIDIKAN KHUSUS BAHASA INDONESIA', 'Membaca dan Memirsa',
     'D', 35, False),
]


@pytest.fixture
def pipeline(test_db):
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


def _params(row, topic='Topik uji'):
    system, inst, grade, subj, elem, _ph, _jp, _kbc = row
    return {'education_system': system, 'institution_type': inst,
            'grade': grade, 'subject': subj, 'element': elem,
            'topic': topic, 'requested_tp_count': 3,
            'jumlah_pertemuan': 2, 'jp_per_pertemuan': 2}


def _rows_for(budget):
    a = budget - 30  # 10 + 20 + a == budget per pertemuan
    return [(1, 'Pembuka', 10, 'TP-1', 'memahami'),
            (1, 'Inti', 20, 'TP-1', 'mengaplikasi'),
            (1, 'Inti', a, 'TP-2', 'mengaplikasi'),
            (2, 'Pembuka', 10, 'TP-2', 'memahami'),
            (2, 'Inti', 20, 'TP-3', 'mengaplikasi'),
            (2, 'Inti', a, 'TP-3', 'merefleksi')]


def _generate(pipeline, row):
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai, mock_meeting_principles)
    budget = row[6] * 2

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            return {'activities': [
                {'name': f'A{i + 1}', 'description': 'Deskripsi ' + f'A{i + 1}',
                 'duration': dur, 'experience': exp, 'tp_linked': tp,
                 'stage': stage, 'meeting': meeting}
                for i, (meeting, stage, dur, tp, exp)
                in enumerate(_rows_for(budget))],
                'meeting_principles': mock_meeting_principles(2)}
        return fake_router_response(prompt, system_instruction, **kwargs)

    with patch_pipeline_ai(pipeline, side_effect):
        return pipeline.generate(_params(row))


def _docx_text(module_dict):
    import tempfile
    import os
    from docx import Document
    from docx_renderer import render_final_rpm_docx
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    try:
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_dict,
             'validation': module_dict.get('validation_results', {})},
            path)
        doc = Document(out)
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for r in table.rows:
                texts.extend(c.text for c in r.cells)
        return '\n'.join(texts)
    finally:
        os.unlink(path)


class TestResolutionMatrix:
    @pytest.mark.parametrize('row', MATRIX,
                             ids=[f'{r[1]}-{r[5]}' for r in MATRIX])
    def test_grade_phase_subject_element_cp(self, pipeline, row):
        system, inst, grade, subj, elem, phase, _jp, _kbc = row
        ok, ctx, errs = pipeline.curriculum_validator.validate(
            _params(row))
        assert ok, errs
        assert ctx.phase == phase
        assert ctx.cp.text and ctx.cp.source_document_id
        assert ctx.cp.source_page

    @pytest.mark.parametrize('row', MATRIX,
                             ids=[f'{r[1]}-{r[5]}' for r in MATRIX])
    def test_cp_routing_by_system(self, pipeline, row):
        system, _inst, _grade, _subj, _elem, _ph, _jp, _kbc = row
        ok, ctx, errs = pipeline.curriculum_validator.validate(
            _params(row))
        assert ok, errs
        doc = ctx.cp.source_document_id
        if system == 'KEMENAG':
            assert '9941' in doc, doc
        else:
            assert '046' in doc or '020' in doc, doc


class TestJpMatrix:
    @pytest.mark.parametrize('row', MATRIX,
                             ids=[f'{r[1]}-{r[5]}' for r in MATRIX])
    def test_jp_matches_matrix(self, row):
        from module_generation_pipeline import resolve_jp_menit
        _sys, inst, _grade, _subj, _elem, phase, expected, _kbc = row
        minutes, err = resolve_jp_menit(inst, phase)
        assert err is None
        assert minutes == expected

    def test_slb_never_falls_back_to_reguler(self):
        from module_generation_pipeline import resolve_jp_menit
        assert resolve_jp_menit('SMALB', 'E') == (40, None)
        assert resolve_jp_menit('SMA', 'E') == (45, None)
        assert resolve_jp_menit('SDLB', 'B') == (30, None)


class TestInvalidUnsupported:
    def test_unknown_institution_fail_closed(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('MILB', 'A')
        assert minutes is None and err

    def test_tklb_fondasi_fail_closed(self):
        from module_generation_pipeline import resolve_jp_menit
        minutes, err = resolve_jp_menit('TKLB', 'Fondasi')
        assert minutes is None and err

    def test_wrong_subject_system_rejected(self, pipeline):
        ok, _ctx, errs = pipeline.curriculum_validator.validate({
            'education_system': 'KEMENAG', 'institution_type': 'MTs',
            'grade': 'MTs_7', 'subject': 'MATEMATIKA',
            'element': 'Aljabar', 'topic': 'x'})
        assert not ok

    def test_element_without_cp_rejected(self, pipeline):
        ok, _ctx, errs = pipeline.curriculum_validator.validate({
            'education_system': 'KEMENDIKDASMEN', 'institution_type': 'SMP',
            'grade': 'SMP_7', 'subject': 'MATEMATIKA',
            'element': 'Kalkulus', 'topic': 'x'})
        assert not ok
        assert any('ELEMENT_NOT_FOUND' in e for e in errs)

    def test_grade_phase_mismatch_rejected(self, pipeline):
        ok, _ctx, errs = pipeline.curriculum_validator.validate({
            'education_system': 'KEMENDIKDASMEN', 'institution_type': 'SMA',
            'grade': 'SMA_99', 'subject': 'MATEMATIKA',
            'element': 'Aljabar dan Fungsi', 'topic': 'x'})
        assert not ok


class TestTemplatePerInstitution:
    @pytest.mark.parametrize('row', MATRIX,
                             ids=[f'{r[1]}-{r[5]}' for r in MATRIX])
    def test_full_template_a_to_e(self, pipeline, row):
        _sys, _inst, _grade, _subj, _elem, phase, jp, kbc = row
        result = _generate(pipeline, row)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        assert module['phase'] == phase
        assert len(module['meetings']) == 2
        for mtg in module['meetings']:
            assert mtg['minutes'] == mtg['used_minutes'] == jp * 2
        text = _docx_text(module)
        for section in ('A. IDENTITAS MODUL', 'B. IDENTIFIKASI',
                        'C. DESAIN PEMBELAJARAN',
                        'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN',
                        'E. ASESMEN'):
            assert section in text, section
        assert 'Provenance CP:' not in text
        assert 'source_fragment_id' not in text
        dims = module['profile_dimensions']
        assert 1 <= len(dims) <= 7
        if kbc:
            assert 'Kurikulum Berbasis Cinta' in text
        else:
            assert 'Kurikulum Berbasis Cinta' not in text
            assert 'Panca Cinta' not in text
