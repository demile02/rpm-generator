#!/usr/bin/env python3
"""Batch 3: pedagogical quality, content quality & usability regressions.

Temuan audit baseline (Phase 0, bukti di audit/batch*_e2e*.json):
- QUALITY_GAP: DOCX merender provenance internal mentah (fragment id,
  kode versi) kepada guru -> disajikan sebagai sitasi manusiawi.
- QUALITY_GAP (kosmetik normatif): error budget tanpa kode
  TIME_BUDGET_INVALID -> semua error validator memakai kode.
- EXPECTED_BEHAVIOR (dipertahankan + dikunci test): rantai alignment
  CP->TP->aktivitas->asesmen, TP C3+, KKTP konsisten dengan TP,
  budget deterministik, diagnostic prasyarat tanpa TP-link, output
  bebas slop bare-pattern, DOCX tanpa metadata internal.
"""
import re
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from module_generation_pipeline import (
    MasterOutlineValidator, validate_time_budget,
)
from test_factuality_gate import module_from_result

MEET_PARAMS = {
    'education_system': 'KEMENAG',
    'institution_type': 'MTs',
    'grade': 'MTs_7',
    'subject': 'Akidah Akhlak',
    'element': 'Pemahaman Konsep',
    'topic': 'Taubat',
    'requested_tp_count': 3,
    'jumlah_pertemuan': 2,
    'jp_per_pertemuan': 2,
}

ROWS_D = [
    (1, 'Pembuka', 15, 'TP-1', 'memahami'),
    (1, 'Inti', 35, 'TP-1', 'memahami'),
    (1, 'Inti', 30, 'TP-2', 'mengaplikasi'),
    (2, 'Pembuka', 15, 'TP-2', 'mengaplikasi'),
    (2, 'Inti', 35, 'TP-3', 'mengaplikasi'),
    (2, 'Penutup', 30, 'TP-3', 'merefleksi'),
]


ACT_NAMES = [
    'Apersepsi pengalaman taubat',
    'Jelajah konsep taubat nasuha',
    'Sortir langkah taubat yang tepat',
    'Kilasan langkah taubat',
    'Simulasi studi kasus taubat',
    'Refleksi komitmen perbaikan diri',
]

ACT_DESCRIPTIONS = [
    'Siswa menyimak dalil taubat dan berbagi pengalaman umum tentang '
    'alasan kembali kepada Allah dalam suasana aman.',
    'Kelompok menganalisis pengertian, syarat, dan tata cara taubat '
    'nasuha melalui kartu informasi lalu menyusun peta konsep.',
    'Siswa mengurutkan kartu tindakan dari situasi kesalahan dan '
    'mencocokkannya dengan langkah taubat nasuha yang benar.',
    'Siswa mengikuti permainan mencocokkan situasi, langkah taubat, '
    'dan ungkapan doa secara cepat dan menyenangkan.',
    'Kelompok kecil memerankan kasus pelanggaran lalu menunjukkan '
    'proses taubat nasuha melalui dialog dan tindakan perbaikan.',
    'Siswa menulis refleksi pribadi tentang kesalahan yang diperbaiki '
    'dan menutup pembelajaran dengan doa bersama.',
]


def _mock_activities():
    from test_phase_c_module_generation import mock_meeting_principles
    return {'activities': [
        {'name': ACT_NAMES[i], 'description': ACT_DESCRIPTIONS[i],
         'duration': dur, 'experience': exp, 'tp_linked': tp,
         'stage': stage, 'meeting': meeting}
        for i, (meeting, stage, dur, tp, exp) in enumerate(ROWS_D)],
        'meeting_principles': mock_meeting_principles(2)}


@pytest.fixture
def pipeline(test_db):
    from module_generation_pipeline import ModuleGenerationPipeline
    pipe = ModuleGenerationPipeline(db_path=str(test_db))
    yield pipe
    pipe.curriculum_validator.close()


def _mock_generate(pipeline, tp_override=None, kktp_override=None,
                   activities_override=None):
    from test_phase_c_module_generation import (
        fake_router_response, patch_pipeline_ai)

    def side_effect(prompt, system_instruction=None, **kwargs):
        if '"activities"' in prompt:
            if activities_override is not None:
                from test_phase_c_module_generation import (
                    mock_meeting_principles)
                return {'activities': activities_override,
                        'meeting_principles': mock_meeting_principles(2)}
            return _mock_activities()
        if 'tp_list' in prompt and tp_override is not None:
            return dict(tp_override)
        if 'criteria' in prompt and kktp_override is not None:
            return dict(kktp_override(prompt))
        return fake_router_response(prompt, system_instruction, **kwargs)

    with patch_pipeline_ai(pipeline, side_effect):
        return pipeline.generate(dict(MEET_PARAMS))


def _docx_paragraphs(module_dict):
    import io
    import os
    import tempfile
    from docx import Document
    from docx_renderer import render_final_rpm_docx
    fd, path = tempfile.mkstemp(suffix='.docx')
    os.close(fd)
    try:
        out = render_final_rpm_docx(
            {'status': 'success', 'module': module_dict,
             'validation': module_dict.get('validation_results', {})},
            path)
        doc = Document(out)
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                texts.extend(c.text for c in row.cells)
        return '\n'.join(texts)
    finally:
        os.unlink(path)


class TestProvenancePresentation:
    """Guru melihat sitasi manusiawi, bukan metadata internal."""

    def test_no_internal_provenance_in_docx(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_paragraphs(result['module'])
        assert 'Provenance CP:' not in text
        assert 'source_fragment_id' not in text
        assert 'page-47' not in text
        assert 'Sumber:' in text
        assert 'halaman' in text

    def test_source_doc_and_page_rendered(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        text = _docx_paragraphs(result['module'])
        prov = result['module']['master_outline']['design']['cp_provenance']
        assert prov['source_document_id'] in text
        assert str(prov['source_page']) in text

    def test_missing_provenance_renders_no_fabrication(self):
        from docx import Document
        from docx_renderer import _render_desain
        doc = Document()
        _render_desain(doc, {'cp': 'Teks CP.', 'cp_provenance': {}})
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(c.text for t in doc.tables
                                 for r in t.rows for c in r.cells)
        assert 'Teks CP.' in text
        assert 'Sumber:' not in text
        assert 'Provenance' not in text


class TestTimeBudgetCode:
    def test_all_errors_carry_code(self):
        acts = [dict(zip(('id', 'meeting', 'stage', 'duration', 'tp_linked'),
                         ('Ax', 9, 'Istirahat', 'lama', 'TP-9')))]
        errors = validate_time_budget(acts, 2, 80, {'TP-1'})
        assert errors
        assert all(e.startswith('TIME_BUDGET_INVALID: ') for e in errors)

    def test_valid_budget_still_empty(self):
        acts = [dict(zip(('id', 'meeting', 'stage', 'duration', 'tp_linked'),
                         (f'A{i}', m, s, d, t)))
                for i, (m, s, d, t, _e) in enumerate(ROWS_D)]
        assert validate_time_budget(acts, 2, 80, {'TP-1', 'TP-2', 'TP-3'}) == []

    def test_zero_duration_rejected_with_code(self):
        base = dict(id='A1', meeting=1, stage='Inti', duration=0,
                    tp_linked='TP-1')
        errors = validate_time_budget([base], 1, 80, {'TP-1'})
        assert any(e.startswith('TIME_BUDGET_INVALID: ') for e in errors)

    def test_pipeline_rejection_keeps_wrapper(self, pipeline):
        rows = [dict(zip(('name', 'description', 'duration', 'tp_linked',
                                 'experience', 'stage', 'meeting'),
                        (ACT_NAMES[i], ACT_DESCRIPTIONS[i], dur, tp, exp,
                         stage, meeting)))
                for i, (meeting, stage, dur, tp, exp) in enumerate(ROWS_D)]
        rows[0]['duration'] = 5  # rusak: pertemuan 1 hanya 70 dari 80
        result = _mock_generate(pipeline, activities_override=rows)
        assert result['status'] == 'failed'
        assert any('Alokasi waktu' in str(e) for e in result['errors'])
        assert any('TIME_BUDGET_INVALID' in str(e)
                   for e in result['errors'])


class TestTpKktpCases:
    """§2: C2, UNKNOWN, mismatch, tanpa-KKTP, multi-TP."""

    def test_c2_tp_rejected(self, pipeline):
        result = _mock_generate(pipeline, tp_override={
            'tp_list': ['Memahami konsep taubat dalam kehidupan sehari-hari',
                        'Menganalisis dampak taubat terhadap perilaku siswa',
                        'Mengevaluasi praktik taubat yang sesuai tuntunan Islam']})
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_unknown_verb_tp_rejected(self, pipeline):
        result = _mock_generate(pipeline, tp_override={
            'tp_list': ['Berdoa sebelum belajar dengan khusyuk setiap hari',
                        'Menganalisis dampak taubat terhadap perilaku siswa',
                        'Mengevaluasi praktik taubat yang sesuai tuntunan Islam']})
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_unknown_level_direct(self, pipeline):
        ok, level, _msg = pipeline.rule_validator.validate_tp_cognitive_level(
            'Berdoa sebelum belajar dengan khusyuk')
        assert ok is False

    def test_kktp_below_tp_rejected(self, pipeline):
        def kktp(prompt):
            from test_phase_c_module_generation import fake_router_response
            resp = fake_router_response(prompt)
            return {'criteria': resp['criteria'], 'level': 'C3'}

        # TP mock: TP-1 C3, TP-2 C4, TP-3 C5. KKTP C3 seragam berarti
        # TP-2 (C4) diukur di bawah levelnya -> wajib ditolak.
        result = _mock_generate(pipeline, kktp_override=kktp)
        assert result['status'] == 'failed'
        assert result['module'] is None

    def test_tp_without_kktp_detected(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = module_from_result(result)
        module.success_criteria = [
            k for k in module.success_criteria if k.tp_id != 'TP-3']
        ok, errors = MasterOutlineValidator.validate(module)
        assert ok is False
        assert any(e['code'] == 'KKTP_ERROR' and 'TP-3' in e['message']
                   for e in errors)

    def test_multiple_tp_each_with_kktp(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        tp_ids = [t['id'] for t in result['module']['learning_objectives']]
        assert tp_ids == ['TP-1', 'TP-2', 'TP-3']
        kk_map = {k['tp_id']: k['id']
                  for k in result['module']['success_criteria']}
        assert set(kk_map) == set(tp_ids)


class TestAlignmentChain:
    """§1: CP -> TP -> aktivitas -> asesmen dapat ditelusuri."""

    def test_full_chain_traceable(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        module = result['module']
        cp_id = module['curriculum_context']['cp']['id']
        for tp in module['learning_objectives']:
            assert tp['cp_reference'] == cp_id, tp['id']
        tp_ids = {t['id'] for t in module['learning_objectives']}
        acts = module['learning_activities']
        for tp_id in tp_ids:
            assert any(a.get('tp_linked') == tp_id for a in acts), tp_id
        linked = set()
        asm = module['assessments'] or {}
        for bucket in ('diagnostic', 'formative'):
            for item in asm.get(bucket) or []:
                if isinstance(item, dict) and item.get('tp_linked'):
                    linked.add(item['tp_linked'])
        for item in (asm.get('summative') or {}).get('items') or []:
            if isinstance(item, dict) and item.get('tp_linked'):
                linked.add(item['tp_linked'])
        for tp_id in tp_ids:
            assert tp_id in linked, tp_id

    def test_experience_stage_independent(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        pairs = {(a.get('stage'), a.get('experience'))
                 for a in result['module']['learning_activities']}
        # Tahap dan pengalaman tidak berpasangan 1:1 (mis. Inti dapat
        # membawa memahami; Pembuka dapat membawa mengaplikasi).
        stages_of_memahami = {s for s, e in pairs if e == 'memahami'}
        assert stages_of_memahami != {'Pembuka'}
        assert ('Pembuka', 'mengaplikasi') in pairs


class TestAntiSlopBaseline:
    """§9: tidak ada pola generik telanjang pada output (mock baseline)."""

    BARE_PATTERNS = [
        r'diharapkan mampu memahami',
        r'\bberdiskusi\b.{0,10}$',
        r'\bmenganalisis\b.{0,8}$',
        r'\bmempresentasikan\b.{0,8}$',
        r'apa yang saya pelajari hari ini',
    ]

    def _ai_texts(self, module):
        texts = []
        for act in module.get('learning_activities') or []:
            texts.append(act.get('name', ''))
            texts.append(act.get('description', ''))
        mat = ((module.get('master_outline') or {}).get('identification')
               or {}).get('material_characteristics') or {}
        texts.append(mat.get('summary', ''))
        for sheet in module.get('lkpd') or []:
            texts.append(sheet.get('instructions', ''))
            texts.append(sheet.get('task', ''))
        return [t for t in texts if t]

    def test_no_bare_generic_patterns(self, pipeline):
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        for text in self._ai_texts(result['module']):
            for pat in self.BARE_PATTERNS:
                assert not re.search(pat, text, re.IGNORECASE | re.MULTILINE), (
                    pat, text[:100])

    def test_activities_name_objects(self, pipeline):
        """Aktivitas menyebut objek/konteks, bukan kata kerja telanjang."""
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        for act in result['module']['learning_activities']:
            assert len((act.get('name') or '').split()) >= 2, act.get('id')
            assert len((act.get('description') or '')) >= 20, act.get('id')


class TestDocxCleanliness:
    """§8: tidak ada raw AI/debug/metadata internal di DOCX."""

    INTERNAL_MARKERS = ['validation_results', 'generation_id', 'prompt_hash',
                        'traceback', 'Exception', 'layer tambahan', '```',
                        'A3_EMPTY', 'TIME_BUDGET_INVALID', 'Provenance CP:',
                        'source_fragment_id']

    def test_docx_free_of_internal_markers(self, pipeline):
        import io
        import os
        import tempfile
        from docx import Document
        from docx_renderer import render_final_rpm_docx
        result = _mock_generate(pipeline)
        assert result['status'] == 'success', result['errors']
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            out = render_final_rpm_docx(
                {'status': 'success', 'module': result['module'],
                 'validation': result['module'].get('validation_results', {})},
                path)
            doc = Document(out)
            texts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    texts.extend(c.text for c in row.cells)
            full = '\n'.join(texts)
        finally:
            os.unlink(path)
        for marker in self.INTERNAL_MARKERS:
            assert marker not in full, marker
