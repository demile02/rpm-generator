#!/usr/bin/env python3
"""Batch 6: export DOCX selalu memakai RPM terbaru/aktif di Web.

Alur: Generate -> Web -> Edit -> Simpan versi baru -> Export.
Export default (tanpa ?version=) wajib berisi versi terbaru, bukan
generation awal. Renderer pasif: hanya merender versi yang diberikan
caller (dibuktikan: ?version=1 tetap merender v1).
"""
import io
import json
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from test_versions_batch5 import (  # noqa: E402
    _save, _seed, api_pipeline,  # noqa: F401 (fixture reuse)
)

TP_BARU = 'Menerapkan taubat setiap hari VERSI GURU'


@pytest.fixture
def seeded(flask_client, api_pipeline):
    gid, module = _seed(flask_client, api_pipeline)
    return gid, module


def _export(flask_client, gid, version=None):
    url = f'/api/module/{gid}/export-docx'
    if version is not None:
        url += f'?version={version}'
    return flask_client.get(url)


def _docx_text(payload: bytes) -> str:
    doc = Document(io.BytesIO(payload))
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return '\n'.join(texts)


def _version_module(flask_client, gid, version_no):
    flask_client.get(f'/api/module/{gid}/versions')
    resp = flask_client.get(f'/api/module/{gid}/versions/{version_no}')
    assert resp.status_code == 200, resp.get_data(as_text=True)
    return json.loads(resp.data)['version']['module']


def _save_tp(flask_client, gid, base, new_text, module=None):
    if module is None:
        module = _version_module(flask_client, gid, base)
    tp = [{'id': t['id'], 'text': t['text']}
          for t in module['learning_objectives']]
    tp[0]['text'] = new_text
    resp = _save(flask_client, gid, base, {'tp': tp})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    return json.loads(resp.data)['version']


class TestExportLatest:
    def test_export_awal_isi_web(self, flask_client, seeded):
        """Export versi awal = isi Web (v1)."""
        gid, module = seeded
        tp1 = module['learning_objectives'][0]['text']
        resp = _export(flask_client, gid)
        assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
        text = _docx_text(resp.data)
        assert tp1[:40] in text

    def test_export_setelah_save_versi_baru(self, flask_client, seeded):
        """Setelah Simpan v2, export default wajib berisi TP-1 = B."""
        gid, module = seeded
        tp1_lama = module['learning_objectives'][0]['text']
        assert tp1_lama != TP_BARU
        saved = _save_tp(flask_client, gid, 1, TP_BARU, module)
        assert saved['version_no'] == 2
        resp = _export(flask_client, gid)
        assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
        text = _docx_text(resp.data)
        assert TP_BARU[:40] in text
        assert tp1_lama not in text

    def test_export_tidak_mengambil_generation_awal(self, flask_client,
                                                    seeded):
        """Tanpa ?version pun export tidak kembali ke v1."""
        gid, module = seeded
        tp1_lama = module['learning_objectives'][0]['text']
        _save_tp(flask_client, gid, 1, TP_BARU, module)
        v2 = _version_module(flask_client, gid, 2)
        _save_tp(flask_client, gid, 2, TP_BARU + ' V3', v2)
        resp = _export(flask_client, gid)
        assert resp.status_code == 200
        text = _docx_text(resp.data)
        assert (TP_BARU + ' V3')[:40] in text
        assert tp1_lama not in text

    def test_export_version_eksplisit_masih_dilayani(self, flask_client,
                                                     seeded):
        """Renderer pasif: ?version=1 merender v1 (bukti caller-driven)."""
        gid, module = seeded
        tp1_lama = module['learning_objectives'][0]['text']
        _save_tp(flask_client, gid, 1, TP_BARU, module)
        resp = _export(flask_client, gid, version=1)
        assert resp.status_code == 200, resp.get_data(as_text=True)[:300]
        text = _docx_text(resp.data)
        assert tp1_lama[:40] in text
        assert TP_BARU not in text


class TestViewSaveCycle:
    """Batch 9: load -> lihat -> simpan tanpa edit tidak boleh error;
    tp.text selalu non-kosong; validator tetap utuh."""

    def test_resave_identik_berhasil(self, flask_client, seeded):
        """Case A (server): simpan ulang TP identik -> v2 identik."""
        import json as _json
        gid, module = seeded
        tp = [{'id': t['id'], 'text': t['text']}
              for t in module['learning_objectives']]
        resp = _save(flask_client, gid, 1, {'tp': tp})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        v2 = _version_module(flask_client, gid, 2)
        for a, b in zip(module['learning_objectives'],
                         v2['learning_objectives']):
            assert a['text'] == b['text']
            assert isinstance(b['text'], str) and b['text'].strip()

    def test_edit_hanya_ubah_target(self, flask_client, seeded):
        """Case B: edit satu TP -> v2 hanya berubah di TP itu."""
        gid, module = seeded
        tp = [{'id': t['id'], 'text': t['text']}
              for t in module['learning_objectives']]
        tp[1]['text'] = 'Menerapkan TP dua yang diedit guru'
        resp = _save(flask_client, gid, 1, {'tp': tp})
        assert resp.status_code == 200, resp.get_data(as_text=True)
        v2 = _version_module(flask_client, gid, 2)
        assert v2['learning_objectives'][1]['text'] == \
            'Menerapkan TP dua yang diedit guru'
        assert v2['learning_objectives'][0]['text'] == \
            module['learning_objectives'][0]['text']
        assert v2['learning_objectives'][2]['text'] == \
            module['learning_objectives'][2]['text']

    def test_tp_kosong_ditolak_validator(self, flask_client, seeded):
        """Validator utuh (aturan 4): tp.text kosong -> 400, tanpa
        version baru."""
        gid, module = seeded
        tp = [{'id': t['id'], 'text': t['text']}
              for t in module['learning_objectives']]
        tp[0]['text'] = '   '
        resp = _save(flask_client, gid, 1, {'tp': tp})
        assert resp.status_code == 400, resp.get_data(as_text=True)
        body = json.loads(
            flask_client.get(f'/api/module/{gid}/versions').data)
        assert body['latest_version'] == 1

    def test_export_setelah_view_sama_aktif(self, flask_client, seeded):
        """Case C: lihat v2 lalu export -> sama dengan versi aktif."""
        gid, module = seeded
        tp = [{'id': t['id'], 'text': t['text']}
              for t in module['learning_objectives']]
        tp[0]['text'] = TP_BARU
        _save(flask_client, gid, 1, {'tp': tp})
        viewed = _version_module(flask_client, gid, 2)
        text = _docx_text(_export(flask_client, gid, version=2).data)
        assert (viewed['learning_objectives'][0]['text'])[:40] in text


class TestLegacyBaseline:
    def test_riwayat_materialisasi_v1(self, flask_client, seeded):
        """Generasi lama tanpa version rows mendapat v1 saat riwayat
        dibuka — tombol regen/simpan/finalize menjadi hidup."""
        import json as _json
        gid, _ = seeded
        resp = flask_client.get(f'/api/module/{gid}/versions')
        assert resp.status_code == 200, resp.get_data(as_text=True)
        body = _json.loads(resp.data)
        assert body['latest_version'] == 1
        assert body['versions'][0]['status'] == 'validated'

    def test_regen_setelah_materialisasi(self, flask_client, seeded,
                                         api_pipeline):
        """Regen asesmen pada generasi lama memakai kontrak Batch 4."""
        import json as _json
        from test_phase_c_module_generation import patch_pipeline_ai
        from test_meetings_rpm import _meetings_fake, ROWS_D
        gid, _ = seeded
        assert flask_client.get(
            f'/api/module/{gid}/versions').status_code == 200
        with patch_pipeline_ai(api_pipeline, _meetings_fake(ROWS_D)):
            resp = flask_client.post(
                f'/api/module/{gid}/regenerate',
                json={'base_version': 1, 'target': 'asesmen'})
        assert resp.status_code == 200, resp.get_data(as_text=True)[:500]
        version = _json.loads(resp.data)['version']
        assert version['version_no'] == 2
        text = _docx_text(_export(flask_client, gid).data)
        assert 'Jawaban Benar' in text
        assert 'multiple_choice' not in text


class TestCurrentVersionParity:
    def test_web_docx_parity_current_version(self, flask_client, seeded):
        """DOCX export v2 memuat field kunci modul v2 (Web menampilkan
        modul yang sama 1:1 — parity Web dibuktikan node tests)."""
        gid, module = seeded
        saved = _save_tp(flask_client, gid, 1, TP_BARU, module)
        assert saved['version_no'] == 2
        current = _version_module(flask_client, gid, 2)
        text = _docx_text(_export(flask_client, gid).data)
        assert (current['topic'] or '').strip() in text
        for tp in current['learning_objectives']:
            assert (tp['text'] or '')[:30] in text
        for mtg in current.get('meetings') or []:
            assert f"Pertemuan {mtg['index']}" in text
        assert 'A. IDENTITAS MODUL' in text
        assert 'E. ASESMEN' in text
        # Regression layout Batch 3/4 pada versi edited.
        assert 'Versi Kurikulum' not in text
        assert 'Total keseluruhan' not in text
        assert 'Tema yang relevan' not in text
        assert 'multiple_choice' not in text
