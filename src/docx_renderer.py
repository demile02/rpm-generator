#!/usr/bin/env python3
"""Renderer DOCX final RPM Pembelajaran Mendalam.

Kontrak (AGENTS.md §12/§32): renderer HANYA menerima hasil FINAL
VALIDATED (``result["status"] == "success"`` dan
``result["validation"]["final"]["passed"] is True``). Selain itu
-> ExportError, tanpa bypass apa pun (tanpa force/skip/fallback).

Single source of truth (AGENTS.md §43):
FINAL VALIDATED RPM RESULT ├──> WEB RENDERER └──> DOCX RENDERER.
Renderer murni representasi: tidak memanggil AI, tidak meregenerasi,
tidak enrich/paraphrase/infer/default, tidak mengubah
CP/TP/level/provenance/validation, tidak menambah konten DOCX-only.

Layout mengikuti ``templates/Templatku.docx`` (AGENTS.md §41/§44,
BATCH 2): A4, margin 2,5 cm, font Bookman Old Style 11 pt, banner
section ter-shade hijau muda (#E2EFD9), border single, banner
merged per section/grup, kolom proporsional template
(C: 11/78/11; D: 18,7/42,2/10,9/9,4/18,8; E: 65,6/12,5/9,4/12,5),
judul terpusat + topik terpusat berkutip, urutan
A. IDENTITAS MODUL → B → C → D (Pertemuan → Pembuka → Inti →
Penutup → Prinsip) → E (Diagnostik/Formatif/Sumatif).

Parity Web (Batch 1): judul, topik, field identitas, B–E, TP/KKTP,
meeting/aktivitas, prinsip per pertemuan, asesmen — sama dengan
Web Preview untuk result yang sama. Glosarium dipertahankan karena
bagian dari final result (ditampilkan Web).
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Geometri Templatku.docx: A4, margin 2,5 cm, banner #E2EFD9.
PAGE_W_CM = 21.0
PAGE_H_CM = 29.7
MARGIN_CM = 2.5
USABLE_CM = PAGE_W_CM - 2 * MARGIN_CM
FONT = 'Bookman Old Style'
BANNER_FILL = 'E2EFD9'
# Border template per section (tblBorders Templatku.docx):
# A hitam (auto), B/C/D biru 365F91, E abu BFBFBF; single sz 4.
BORDER_AUTO = 'auto'
BORDER_SECT = '365F91'
BORDER_ASESMEN = 'BFBFBF'
INK = '000000'
MUTED = '595959'

# Proporsi kolom template (gridCol Templatku.docx).
# E: grid 4 kolom Instrumen|Tipe|TP|KKTP; diagnostik memakai 3 kolom
# pertama dengan 2 terakhir merged (Jawaban Benar).
TP_FRACS = [0.109, 0.782, 0.109]
ACTIVITY_FRACS = [0.187, 0.422, 0.109, 0.094, 0.188]
ASSESS_FRACS = [0.6561, 0.1251, 0.0937, 0.1251]
GLOS_FRACS = [0.25, 0.75]

BODY_SIZE = Pt(11)
SMALL_SIZE = Pt(9)
TITLE_SIZE = Pt(11)
TOPIC_SIZE = Pt(11)
BANNER_SIZE = Pt(12)
SUB_SIZE = Pt(11)


class ExportError(Exception):
    """Export ditolak: input bukan FINAL VALIDATED RPM."""


def assert_final_validated(result) -> dict:
    """Kembalikan dict module bila result FINAL VALIDATED.

    Raise ExportError untuk: bukan dict, status != success, final
    validation tidak PASS, atau module/master_outline hilang.
    """
    if not isinstance(result, dict):
        raise ExportError("Export ditolak: hasil bukan dict result pipeline.")
    if result.get('status') != 'success':
        raise ExportError(
            "Export ditolak: status=%r (final validation tidak lulus atau "
            "draft belum tervalidasi)." % (result.get('status'),))
    validation = result.get('validation') or {}
    final = validation.get('final') or {}
    if final.get('passed') is not True:
        raise ExportError(
            "Export ditolak: final validation != PASS "
            "(errors=%r)." % (final.get('errors'),))
    module = result.get('module')
    if not isinstance(module, dict) or not isinstance(
            module.get('master_outline'), dict):
        raise ExportError(
            "Export ditolak: module/master_outline final tidak tersedia.")
    return module


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Ya' if value else 'Tidak'
    return str(value).strip()


# ------------------------------------------------------------------
# Display mapping presentasi (data final tidak diubah; hanya cara
# tampil mengikuti Templatku.docx).
# ------------------------------------------------------------------

def _display_grade(code) -> str:
    """Kode 'MA_10' -> '10' (presentation-friendly, template §2)."""
    code = _text(code)
    if '_' in code:
        tail = code.split('_', 1)[1].strip()
        if tail:
            return tail.replace('_', ' ')
    return code


def _display_semester(value) -> str:
    """1 -> 'Ganjil', 2 -> 'Genap'; selain itu tampil apa adanya."""
    text = _text(value)
    if text == '1':
        return 'Ganjil'
    if text == '2':
        return 'Genap'
    return text


_TYPE_LABELS = {
    'multiple_choice': 'Pilihan Ganda',
    'short_answer': 'Isian',
    'essay': 'Uraian',
    'true_false': 'Benar/Salah',
    'matching': 'Menjodohkan',
    'performance_task': 'Unjuk Kerja',
    'performance_assessment': 'Unjuk Kerja',
    'performance': 'Unjuk Kerja',
    'oral_reading': 'Membaca Nyaring',
}


def _display_type(value) -> str:
    """Enum tipe -> Bahasa Indonesia (template §8)."""
    text = _text(value)
    return _TYPE_LABELS.get(text, text)


def _answer_label(item: dict, options) -> str:
    """Label A/B/C/D untuk kolom Jawaban Benar diagnostik (Batch 4.1).

    Presentation-only: 'D' -> 'D'; 'D. teks...' -> 'D';
    teks lengkap -> label opsi yang cocok; tak cocok -> nilai mentah
    apa adanya (jujur, tanpa karangan).
    """
    if not isinstance(item, dict):
        return ''
    raw = ''
    for key in ('correct_answer', 'answer_key', 'answer'):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            raw = value.strip()
            break
        if isinstance(value, list) and value:
            raw = _text(value[0])
            break
    if not raw:
        return ''
    labels = [o['label'] for o in (options or [])
              if isinstance(o, dict) and o.get('label')]
    if len(raw) == 1 and raw.upper() in labels:
        return raw.upper()
    first = raw.split('.', 1)[0].strip().upper()
    if len(first) == 1 and first in labels:
        return first
    for opt in options or []:
        if isinstance(opt, dict) and opt.get('text', '').strip().lower() \
                == raw.lower():
            return opt['label']
    return raw


def _mc_options(item: dict):
    """Opsi Pilihan Ganda ternormalisasi [{label, text}] (Batch 4)."""
    options = (item or {}).get('options') if isinstance(item, dict) else None
    if not isinstance(options, list):
        return []
    out = []
    for i, opt in enumerate(options):
        if isinstance(opt, dict):
            label = _text(opt.get('label')) or chr(ord('A') + i)
            text = _text(opt.get('text', opt.get('option', '')))
        else:
            label, text = chr(ord('A') + i), _text(opt)
        if text:
            out.append({'label': label, 'text': text})
    return out


# ------------------------------------------------------------------
# Primitif template (Oxml + paragraf).
# ------------------------------------------------------------------

def _run(paragraph, text: str, bold: bool = False, size=BODY_SIZE,
         italic: bool = False, color_hex: str = INK):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = size
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(color_hex)
    return run


def _para(doc_or_cell, text: str, bold: bool = False, size=BODY_SIZE,
          italic: bool = False, align=None, space_after=Pt(4),
          color_hex: str = INK):
    if hasattr(doc_or_cell, 'sections'):
        paragraph = doc_or_cell.add_paragraph()
    else:
        paras = doc_or_cell.paragraphs
        if len(paras) == 1 and not paras[0].text:
            paragraph = paras[0]
        else:
            paragraph = doc_or_cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = space_after
    if align is not None:
        paragraph.alignment = align
    if text:
        _run(paragraph, text, bold=bold, size=size, italic=italic,
             color_hex=color_hex)
    return paragraph


def _shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def _set_table_borders(table, size: str = '4',
                       color: str = BORDER_AUTO) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)
        borders.append(element)
    tbl_pr.append(borders)


def _dxa(cm_value: float) -> int:
    return int(round(cm_value * 567.0))  # 1 cm = 567 dxa


def _set_tc_width(tc, width_dxa: int) -> None:
    tc_pr = tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:tcW')):
        tc_pr.remove(old)
    tc_w = OxmlElement('w:tcW')
    tc_w.set(qn('w:w'), str(width_dxa))
    tc_w.set(qn('w:type'), 'dxa')
    tc_pr.append(tc_w)


def _fix_layout(table, fractions) -> None:
    """Kunci layout fixed + lebar kolom proporsional template.

    Ditulis eksplisit ke tblGrid/tcW (Oxml) karena setter
    ``cell.width`` python-docx tidak akurat pada baris merged.
    """
    table.autofit = False
    try:
        table.allow_autofit = False
    except AttributeError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    usable = _dxa(USABLE_CM)
    widths = [int(round(usable * f)) for f in fractions]
    widths[-1] += usable - sum(widths)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(usable))
    tbl_w.set(qn('w:type'), 'dxa')
    grid = tbl.find(qn('w:tblGrid'))
    if grid is None:
        grid = OxmlElement('w:tblGrid')
        tbl.append(grid)
    for old in grid.findall(qn('w:gridCol')):
        grid.remove(old)
    for width in widths:
        column = OxmlElement('w:gridCol')
        column.set(qn('w:w'), str(width))
        grid.append(column)
    for row in table.rows:
        cursor = 0
        for tc in row._tr.findall(qn('w:tc')):
            tc_pr = tc.get_or_add_tcPr()
            span_el = tc_pr.find(qn('w:gridSpan'))
            merged_away = (tc_pr.find(qn('w:hMerge')) is not None
                           and span_el is None)
            if merged_away:
                if cursor < len(widths):
                    _set_tc_width(tc, widths[cursor])
                cursor += 1
                continue
            span = int(span_el.get(qn('w:val'))) if span_el is not None else 1
            _set_tc_width(tc, sum(widths[cursor:cursor + span]))
            cursor += span
    for row in table.rows:
        for cell in row.cells:
            if not cell._tc.xpath('.//w:vAlign'):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:tblHeader'))


def _center_cell(cell) -> None:
    """Rata tengah isi sel (kolom kode pendek ala Templatku)."""
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER


def _new_table(doc, ncols):
    table = doc.add_table(rows=0, cols=ncols)
    table.style = 'Table Grid'
    return table


def _banner_row(table, text: str) -> None:
    """Baris banner merged + shade ala Templatku.docx."""
    ncols = len(table.columns)
    row = table.add_row()
    row.cells[0].merge(row.cells[ncols - 1])
    cell = row.cells[0]
    _shade_cell(cell, BANNER_FILL)
    _para(cell, _text(text), bold=True, size=BANNER_SIZE,
          space_after=Pt(2))


def _head_row(table, headers) -> None:
    """Baris header kolom: bold tanpa shade (seperti template)."""
    row = table.add_row()
    for i, head in enumerate(headers):
        _para(row.cells[i], _text(head), bold=True, size=BODY_SIZE,
              space_after=Pt(2))
    _repeat_header(row)
    return row


def _body_row(table, values) -> None:
    row = table.add_row()
    for i in range(len(table.columns)):
        _para(row.cells[i],
              _text(values[i]) if i < len(values) else '',
              space_after=Pt(2))
    return row


def _bullets_cell(cell, items, indent: bool = False) -> None:
    for item in items or []:
        text = _text(item.get('text', item) if isinstance(item, dict)
                     else item)
        if not text:
            continue
        paragraph = cell.add_paragraph(style='List Bullet')
        if indent:
            paragraph.paragraph_format.left_indent = Cm(0.75)
        _run(paragraph, text)


# ------------------------------------------------------------------
# Section A–E + Glosarium (struktur Templatku.docx, isi final result).
# ------------------------------------------------------------------

def _render_identitas(doc, identity: dict, module: dict):
    # Template §2: presentation label-value (bukan grid 2 kolom),
    # tanpa Versi Kurikulum, tanpa metadata internal.
    table = _new_table(doc, 1)
    _banner_row(table, 'A. IDENTITAS MODUL')
    ident = module.get('module_identity') or {}
    ctx = module.get('curriculum_context') or {}
    meetings = module.get('meetings') or []
    alokasi = ''
    if meetings:
        total_jp = sum(int(m.get('jp') or 0) for m in meetings
                       if isinstance(m, dict))
        if total_jp > 0:
            jps = {int(m['jp']) for m in meetings
                   if isinstance(m, dict) and int(m.get('jp') or 0) > 0}
            alokasi = (f"{total_jp} JP ({len(meetings)} Pertemuan"
                       + (f" x {sorted(jps)[0]} JP" if len(jps) == 1 else '')
                       + ")")
    kfs = '/'.join(_text(v) for v in
                   (_display_grade(module.get('grade')),
                    _text(module.get('phase')),
                    _display_semester(ident.get('semester'))) if _text(v))
    nama_label = ('Nama Madrasah' if ctx.get('education_system') == 'KEMENAG'
                  else 'Nama Sekolah')
    rows = [
        (nama_label, ident.get('satuan_pendidikan')),
        ('Nama Penyusun', module.get('penyusun')),
        ('Mata Pelajaran', module.get('subject')),
        ('Kelas / Fase /Semester', kfs),
        ('Alokasi Waktu', alokasi),
        ('Tahun Pelajaran', ident.get('year')),
    ]
    row = table.add_row()
    for label, value in rows:
        if _text(value):
            paragraph = row.cells[0].add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5.0))
            _run(paragraph, f"{label}\t: {_text(value)}")
    _fix_layout(table, [1.0])
    _set_table_borders(table, color=BORDER_AUTO)


def _render_identifikasi(doc, identification: dict, module: dict,
                         general: dict):
    # Template §4: satu blok dokumen (satu sel isi), bukan row-table
    # terpisah per subbagian.
    table = _new_table(doc, 1)
    _banner_row(table, 'B. IDENTIFIKASI')
    row = table.add_row()
    cell = row.cells[0]
    readiness = identification.get('learner_readiness') or {}
    _para(cell, 'Kesiapan Murid', bold=True, size=SUB_SIZE,
          space_after=Pt(2))
    _para(cell, _text(readiness.get('summary')
                      or readiness.get('note')
                      or 'Belum ada data kesiapan.'),
          space_after=Pt(2))
    for aspect in readiness.get('aspects') or []:
        paragraph = cell.add_paragraph(style='List Bullet')
        _run(paragraph, _text(aspect))
    material = identification.get('material_characteristics') or {}
    _para(cell, 'Karakteristik Materi', bold=True, size=SUB_SIZE,
          space_after=Pt(2))
    _para(cell,
          _text(material.get('summary') or 'Belum ada ringkasan.'),
          space_after=Pt(2))
    dims = identification.get('profile_dimensions') or []
    dim_notes = identification.get('profile_dimension_notes') or {}
    if dims:
        _para(cell, 'Dimensi Profil Lulusan', bold=True,
              size=SUB_SIZE, space_after=Pt(2))
        for number, dim in enumerate(dims, start=1):
            _para(cell, f"{number}. {_text(dim)}", bold=True,
                  space_after=Pt(1))
            note = dim_notes.get(dim) if isinstance(dim_notes, dict) else None
            if isinstance(note, str) and note.strip():
                _para(cell, _text(note), space_after=Pt(2))
    _render_kbc_di_identifikasi(cell, module, general)
    _fix_layout(table, [1.0])
    _set_table_borders(table, color=BORDER_SECT)


def pair_kbc_items(themes, insertions):
    """Pasangan bernomor (tema, insersi) untuk render KBC (Batch 4).

    Aturan deterministik, tanpa karangan:
    1. insersi ber-field 'tema' dipasangkan ke tema yang sama
       (setiap insersi dipakai sekali);
    2. bila TIDAK ADA info tema sama sekali dan jumlah seimbang,
       pasangan posisional;
    3. selain itu: tema tanpa insersi tampil tanpa rincian, insersi
       tanpa tema tampil sebagai butir bernomor sendiri — tidak ada
       pasangan yang dikarang.
    Kembalikan list [(tema_atau_None, teks_atau_None)].
    """
    themes = [_text(t) for t in (themes or []) if _text(t)]
    items = []
    for ins in insertions or []:
        if isinstance(ins, dict):
            text = _text(ins.get('text'))
            tema = ins.get('tema')
            tema = tema if isinstance(tema, str) and tema.strip() else None
        else:
            text, tema = _text(ins), None
        if text:
            items.append({'text': text, 'tema': tema})
    pairs, used = [], set()
    any_tema = any(it['tema'] for it in items)
    for pos, theme in enumerate(themes):
        idx = None
        if any_tema:
            idx = next((i for i, it in enumerate(items)
                        if i not in used and it['tema'] == theme), None)
        elif items and len(items) == len(themes) and pos not in used:
            idx = pos
        if idx is not None:
            used.add(idx)
            pairs.append((theme, items[idx]['text']))
        else:
            pairs.append((theme, None))
    for i, it in enumerate(items):
        if i not in used:
            pairs.append((None, it['text']))
    return pairs


def _render_kbc_di_identifikasi(cell, module: dict, general: dict) -> bool:
    """KBC sebagai blok B. Identifikasi (khusus madrasah)."""
    ctx = module.get('curriculum_context') or {}
    if ctx.get('education_system') != 'KEMENAG':
        return False
    kbc = general.get('kbc') or {}
    themes = kbc.get('themes') or []
    if 'insertions' in kbc:
        insertion = [ins for ins in (kbc.get('insertions') or [])
                     if isinstance(ins, dict) and ins.get('active')
                     and _text(ins.get('text'))]
    else:
        insertion = kbc.get('insertion_material') or []
    if not (themes or insertion):
        return False
    _para(cell, 'Kurikulum Berbasis Cinta & Materi Insersi',
          bold=True, size=SUB_SIZE, space_after=Pt(2))
    # Batch 4: pasangan bernomor tema + insersi terkait (bukan dua
    # blok daftar terpisah).
    for number, (theme, text) in enumerate(
            pair_kbc_items(themes, insertion), start=1):
        if theme:
            _para(cell, f"{number}. {_text(theme)}", bold=True,
                  space_after=Pt(1))
        else:
            _para(cell, f"{number}. Materi insersi", bold=True,
                  space_after=Pt(1))
        if text:
            _para(cell, _text(text), space_after=Pt(2))
    return True


def _render_desain(doc, design: dict):
    table = _new_table(doc, 3)
    _banner_row(table, 'C. DESAIN PEMBELAJARAN')
    row = table.add_row()
    row.cells[0].merge(row.cells[2])
    _para(row.cells[0], 'Capaian Pembelajaran (CP)', bold=True,
          size=SUB_SIZE, space_after=Pt(2))
    _para(row.cells[0], _text(design.get('cp') or 'CP tidak tersedia.'),
          italic=True, space_after=Pt(2))
    prov = design.get('cp_provenance') or {}
    source_doc = _text(prov.get('source_document_id'))
    if source_doc:
        line = f"Sumber: {source_doc}"
        page = prov.get('source_page')
        if isinstance(page, int) or (
                isinstance(page, str) and page.strip()):
            line += f", halaman {_text(page)}"
        _para(row.cells[0], line, size=SMALL_SIZE, color_hex=MUTED,
              space_after=Pt(2))
    row = table.add_row()
    row.cells[0].merge(row.cells[2])
    _para(row.cells[0], 'Tujuan Pembelajaran', bold=True, size=SUB_SIZE,
          space_after=Pt(2))
    _head_row(table, ['Nomor', 'Tujuan Pembelajaran', 'Level'])
    for t in design.get('tp') or []:
        row = _body_row(table, [_text(t.get('id')), _text(t.get('text')),
                                _text(t.get('cognitive_level'))])
        _center_cell(row.cells[0])
        _center_cell(row.cells[2])
    kktp = design.get('kktp') or []
    if kktp:
        row = table.add_row()
        row.cells[0].merge(row.cells[2])
        _para(row.cells[0], 'Kriteria Ketercapaian (KKTP)', bold=True,
              size=SUB_SIZE, space_after=Pt(2))
        for number, entry in enumerate(kktp, start=1):
            _para(row.cells[0], f"{number}. {_text(entry.get('id'))}",
                  bold=True, space_after=Pt(1))
            for crit in entry.get('criteria') or []:
                _para(row.cells[0], _text(crit), space_after=Pt(1))
    # Template §5: Topik TIDAK dirender di C (sudah di bawah judul).
    practices = design.get('pedagogical_practices') or []
    if practices:
        row = table.add_row()
        row.cells[0].merge(row.cells[2])
        _para(row.cells[0], 'Praktik Pedagogis', bold=True, size=SUB_SIZE,
              space_after=Pt(2))
        for number, practice in enumerate(practices, start=1):
            _para(row.cells[0], f"{number}. {_text(practice)}",
                  space_after=Pt(1))
    _fix_layout(table, TP_FRACS)
    _set_table_borders(table, color=BORDER_SECT)


def _render_pengalaman(doc, learning_experience: dict, meetings) -> None:
    meetings = meetings or []
    table = _new_table(doc, 5)
    _banner_row(table, 'D. LANGKAH DAN PENGALAMAN PEMBELAJARAN')
    if meetings:
        for meeting in meetings:
            if not isinstance(meeting, dict):
                continue
            index = meeting.get('index')
            minutes = meeting.get('minutes') or 0
            row = table.add_row()
            row.cells[0].merge(row.cells[4])
            _para(row.cells[0],
                  f"Pertemuan {index}: {meeting.get('jp')} JP "
                  f"({minutes} menit)", bold=True, size=SUB_SIZE,
                  space_after=Pt(2))
            stages = meeting.get('stages') or {}
            for stage in ('Pembuka', 'Inti', 'Penutup'):
                acts = stages.get(stage) or []
                row = table.add_row()
                row.cells[0].merge(row.cells[4])
                # Template §6: tanpa shading pada tahap; shade
                # hanya untuk banner section.
                _para(row.cells[0], stage, bold=True, size=SUB_SIZE,
                      space_after=Pt(2))
                if acts:
                    _head_row(table, ['Aktivitas', 'Deskripsi', 'Durasi',
                                      'TP', 'Pengalaman'])
                    for a in acts:
                        if not isinstance(a, dict):
                            continue
                        row = _body_row(table, [
                            _text(a.get('name')),
                            _text(a.get('description')),
                            _text(a.get('duration')),
                            _text(a.get('tp_linked')
                                  or a.get('tp_reference')),
                            _text(a.get('experience'))])
                        _center_cell(row.cells[2])
                        _center_cell(row.cells[3])
                        _center_cell(row.cells[4])
            # Templatku.docx: blok "Prinsip Pembelajaran Mendalam"
            # menutup tiap pertemuan (setelah Penutup), bukan blok
            # tersendiri di bawah D. Satu merged-row per pertemuan
            # yang memiliki prinsip.
            principles = meeting.get('principles') or {}
            prinsip_items = []
            for _label, _key in (
                    ('Berkesadaran (Mindful Learning)', 'berkesadaran'),
                    ('Bermakna (Meaningful Learning)', 'bermakna'),
                    ('Menggembirakan (Joyful Learning)',
                     'menggembirakan')):
                _text_pr = _text(principles.get(_key)) \
                    if isinstance(principles, dict) else ''
                if _text_pr:
                    prinsip_items.append((_label, _text_pr))
            if prinsip_items:
                row = table.add_row()
                row.cells[0].merge(row.cells[4])
                _para(row.cells[0], 'Prinsip Pembelajaran Mendalam',
                      bold=True, size=SUB_SIZE, space_after=Pt(2))
                for _number, (_label, _text_pr) in enumerate(
                        prinsip_items, start=1):
                    _para(row.cells[0], f"{_number}. {_label}",
                          bold=True, space_after=Pt(1))
                    _para(row.cells[0], _text_pr, space_after=Pt(2))
            # Template §6: tanpa baris Total / Total keseluruhan.
        _fix_layout(table, ACTIVITY_FRACS)
        _set_table_borders(table, color=BORDER_SECT)
        return
    experiences = (learning_experience.get('experiences') or {})
    order = [k for k in ('memahami', 'mengaplikasi', 'merefleksi')
             if k in experiences] + [
                 k for k in experiences
                 if k not in ('memahami', 'mengaplikasi', 'merefleksi')]
    for exp in order:
        acts = experiences.get(exp) or []
        row = table.add_row()
        row.cells[0].merge(row.cells[4])
        _para(row.cells[0], f"Pengalaman {exp}", bold=True,
              size=SUB_SIZE, space_after=Pt(2))
        for a in acts:
            if not isinstance(a, dict):
                continue
            row = _body_row(table, [
                _text(a.get('name')), _text(a.get('description')),
                _text(a.get('duration')),
                _text(a.get('tp_linked') or a.get('tp_reference')), ''])
            _center_cell(row.cells[2])
            _center_cell(row.cells[3])
    _fix_layout(table, ACTIVITY_FRACS)
    _set_table_borders(table, color=BORDER_SECT)


def _render_asesmen(doc, assessment: dict):
    # Template §8: Diagnostik 3 kolom + Jawaban Benar; Formatif dan
    # Sumatif 4 kolom + TP/KKTP. Tipe dalam Bahasa Indonesia.
    table = _new_table(doc, 4)
    _banner_row(table, 'E. ASESMEN')
    for label, key, diagnostic in (('Asesmen Diagnostik', 'awal', True),
                                   ('Asesmen Formatif', 'proses', False),
                                   ('Asesmen Sumatif', 'akhir', False)):
        items = assessment.get(key)
        row = table.add_row()
        row.cells[0].merge(row.cells[3])
        _shade_cell(row.cells[0], BANNER_FILL)
        _para(row.cells[0], label, bold=True, size=SUB_SIZE,
              space_after=Pt(2))
        if isinstance(items, dict):
            items = items.get('items', [])
        if diagnostic:
            rows = []
            for i in (items or []):
                if not isinstance(i, dict):
                    continue
                opts = _mc_options(i)
                rows.append((_text(i.get('question')),
                             _display_type(i.get('type')),
                             _answer_label(i, opts), opts))
            if rows:
                # Header template: Instrumen | Tipe | Jawaban Benar
                # (merged 2 kolom terakhir, grid 4 kolom dipertahankan).
                row = table.add_row()
                _para(row.cells[0], 'Instrumen', bold=True,
                      space_after=Pt(2))
                _para(row.cells[1], 'Tipe', bold=True, space_after=Pt(2))
                row.cells[2].merge(row.cells[3])
                _para(row.cells[2], 'Jawaban Benar', bold=True,
                      space_after=Pt(2))
                _repeat_header(row)
                for values in rows:
                    row = table.add_row()
                    _para(row.cells[0], values[0], space_after=Pt(2))
                    for opt in values[3]:
                        _para(row.cells[0],
                              f"{opt['label']}. {opt['text']}",
                              space_after=Pt(1))
                    _para(row.cells[1], values[1], space_after=Pt(2))
                    _center_cell(row.cells[1])
                    row.cells[1].vertical_alignment = \
                        WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    row.cells[2].merge(row.cells[3])
                    _para(row.cells[2], values[2], space_after=Pt(2))
                    _center_cell(row.cells[2])
                    row.cells[2].vertical_alignment = \
                        WD_CELL_VERTICAL_ALIGNMENT.CENTER
            elif items:
                row = table.add_row()
                row.cells[0].merge(row.cells[3])
                _bullets_cell(row.cells[0], items)
            continue
        rows = [(_text(i.get('question')), _display_type(i.get('type')),
                 _text(i.get('tp_linked')), _text(i.get('kktp_linked')))
                for i in (items or []) if isinstance(i, dict)]
        if rows:
            _head_row(table, ['Instrumen', 'Tipe', 'TP', 'KKTP'])
            for values in rows:
                row = _body_row(table, list(values))
                _center_cell(row.cells[1])
                _center_cell(row.cells[2])
                _center_cell(row.cells[3])
        elif items:
            row = table.add_row()
            row.cells[0].merge(row.cells[3])
            _bullets_cell(row.cells[0], items)
    _fix_layout(table, ASSESS_FRACS)
    _set_table_borders(table, color=BORDER_ASESMEN)


def _render_lampiran(doc, outline: dict):
    attachments = outline.get('attachments') or {}
    glosarium = attachments.get('glosarium') or []
    if not glosarium:
        return False
    table = _new_table(doc, 2)
    _banner_row(table, 'Glosarium')
    _head_row(table, ['Istilah', 'Definisi'])
    for g in glosarium:
        if isinstance(g, dict) and _text(g.get('istilah')):
            _body_row(table, [_text(g.get('istilah')),
                              _text(g.get('definisi'))])
    _fix_layout(table, GLOS_FRACS)
    _set_table_borders(table)
    return True


def _apply_base_styles(doc) -> None:
    normal = doc.styles['Normal']
    normal.font.size = BODY_SIZE
    normal.font.name = FONT
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(4)
    for name, size in (('Heading 1', Pt(14)), ('Heading 2', Pt(12)),
                       ('Heading 3', Pt(11))):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.size = size
        style.font.name = FONT
        style.font.color.rgb = RGBColor.from_string(INK)
        style.font.bold = True


def _template_status() -> dict:
    """Status template master ``templates/Templatku.docx``.

    Renderer membangun DOCX dari kode (fallback eksplisit) agar export
    tidak bergantung pada file template yang bisa hilang/berubah diam-diam.
    Fungsi ini memberi sinyal jujur: template ADA/HILANG + siapa sumber
    layout aktif. Dipakai health-check dan test template gate.
    """
    from pathlib import Path as _Path
    tpl = _Path(__file__).resolve().parents[1] / 'templates' / 'Templatku.docx'
    return {
        'exists': tpl.is_file(),
        'path': str(tpl),
        'layout_source': ('Templatku.docx (acuan visual)'
                          if tpl.is_file() else 'fallback kode (template hilang)'),
    }


def render_final_rpm_docx(result: dict, output_path) -> str:
    """Render FINAL VALIDATED RPM menjadi file .docx.

    Raise ExportError bila input bukan hasil final tervalidasi.
    Tidak memanggil AI; tidak mengubah data (deepcopy untuk
    keamanan). Return path file sebagai string.
    """
    module = assert_final_validated(result)
    module = deepcopy(module)
    outline = module['master_outline']

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(PAGE_W_CM)
    section.page_height = Cm(PAGE_H_CM)
    section.top_margin = Cm(MARGIN_CM)
    section.bottom_margin = Cm(MARGIN_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)
    _apply_base_styles(doc)

    _para(doc, 'PERENCANAAN PEMBELAJARAN MENDALAM', bold=True,
          size=TITLE_SIZE, align=WD_PARAGRAPH_ALIGNMENT.CENTER,
          space_after=Pt(6))
    topik = (module.get('topic') or (module.get('unit') or {}).get('name')
             or '')
    if topik.strip():
        _para(doc, f"\u201c{topik.strip()}\u201d", size=TOPIC_SIZE,
              align=WD_PARAGRAPH_ALIGNMENT.CENTER, space_after=Pt(12))

    general = outline.get('general_information') or {}
    _render_identitas(doc, general.get('identity') or {}, module)
    _render_identifikasi(doc, outline.get('identification') or {},
                         module, general)
    _render_desain(doc, outline.get('design') or {})
    _render_pengalaman(doc, outline.get('learning_experience') or {},
                       module.get('meetings'))
    _render_asesmen(doc, outline.get('assessment') or {})
    _render_lampiran(doc, outline)

    # Template tidak memiliki footer metadata; ID/generation tetap
    # tertelusur via JSON/DB/URL, bukan badan dokumen (§3).
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.suffix.lower() != '.docx':
        out = out.with_suffix('.docx')
    doc.save(str(out))
    return str(out)
